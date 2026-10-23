"""Named AFC RAN profiles, individual start dates and a real Excel export."""
import datetime as dt
import hashlib
import hmac
import io
import json
import secrets

from flask import abort, jsonify, render_template, request, send_file, session
from werkzeug.utils import secure_filename

import afc_enrolments
import afc_schedule
import afc_positioning

PROFILES = (54, 40, 0)
COMMON_HOURS = (273, 21, 45)
# Dates read from the user's AFC planning, issued on 23 October 2026.
# Scope this initial configuration to that session; never reuse it for another cohort.
PLANNING_SESSION = "1c391db936"
PLANNING_DATES = {"54": "2026-11-16", "40": "2026-11-18", "0": "2026-11-26"}
TARGET = {"54": 8, "40": 2, "0": 2}


def config(training):
    saved = training.get("afc_hours") or {}
    initial = (PLANNING_DATES if training.get("id") == PLANNING_SESSION
               else {"54": training.get("date_start") or "", "40": "", "0": ""})
    return {"dates": {key: (saved.get("dates") or {}).get(key, value)
                      for key, value in initial.items()},
            "target": dict(TARGET) if training.get("id") == PLANNING_SESSION else None}


def participants(services, training):
    return sorted(services._registered_trainees(training), key=services._trainee_alpha_sort_key)


def suggestion(services, training, data=None):
    return afc_positioning.allocation(participants(services, training),
                                     (data or {}).get("positioning_tests") or [], config(training)["target"])


def revision(services, training, data=None):
    value = {"config": config(training), "start": training.get("date_start"),
             "end": training.get("date_end"), "scores": suggestion(services, training, data)["scores"], "trainees": [
                 {key: trainee.get(key) for key in (
                     "id", "first_name", "last_name", "afc_ran_hours", "training_date_start")}
                 for trainee in participants(services, training)]}
    return hashlib.sha256(json.dumps(value, sort_keys=True).encode()).hexdigest()


def view(services, training, data=None):
    settings = config(training)
    proposed = suggestion(services, training, data)
    rows = []
    counts = {str(hours): 0 for hours in PROFILES}
    for trainee in participants(services, training):
        hours = trainee.get("afc_ran_hours")
        if type(hours) is not int or hours not in PROFILES:
            hours = None
        start = settings["dates"].get(str(hours), "")
        actual_start = afc_schedule.effective_session(training, trainee).get("date_start", "")
        if hours is not None:
            counts[str(hours)] += 1
        rows.append({"id": str(trainee["id"]),
                     "last_name": services.normalize_last_name(trainee.get("last_name") or ""),
                     "first_name": services.normalize_first_name(trainee.get("first_name") or ""),
                     "ran_hours": hours, "start": start, "current_start": actual_start,
                     "test": proposed["scores"].get(str(trainee["id"])),
                     "suggested_hours": (proposed["profiles"] or {}).get(str(trainee["id"])),
                     "date_conflict": hours is not None and start != actual_start})
    return {**settings, "rows": rows, "counts": counts, "suggestion": proposed,
            "revision": revision(services, training, data),
            "csrf": session.setdefault("afc_hours_csrf", secrets.token_urlsafe(32)),
            "read_only": session.get("admin_role") == "viewer"}


def validate_dates(training, dates):
    if not isinstance(dates, dict) or set(dates) != {"54", "40", "0"}:
        raise ValueError("Renseignez les dates d’entrée des trois profils RAN.")
    cleaned = {str(hours): afc_schedule.date(dates[str(hours)], f"Entrée RAN {hours} h")
               for hours in PROFILES}
    for value in cleaned.values():
        if value and ((training.get("date_start") and value < training["date_start"])
                      or (training.get("date_end") and value > training["date_end"])):
            raise ValueError("Les dates d’entrée doivent être comprises dans la période de la session.")
    ordered = [cleaned[str(hours)] for hours in PROFILES if cleaned[str(hours)]]
    if ordered != sorted(ordered) or len(ordered) != len(set(ordered)):
        raise ValueError("L’entrée à 54 h doit précéder celle à 40 h, puis celle sans RAN.")
    return cleaned


def export_error(state):
    if not state["rows"]:
        return "Aucun stagiaire inscrit à exporter."
    if any(row["ran_hours"] is None for row in state["rows"]):
        return "Choisissez un profil RAN pour chaque candidat avant l’export."
    if any(not row["start"] for row in state["rows"]):
        return "Renseignez la date de début de chaque profil utilisé."
    if state["target"] and state["counts"] != state["target"]:
        return "La répartition attendue est de 8 candidats à 54 h, 2 à 40 h et 2 sans RAN."
    if any(row["date_conflict"] for row in state["rows"]):
        return "Une date individuelle a changé. Vérifiez les profils puis enregistrez avant l’export."
    return ""


def update(services, data, session_id, payload):
    training = services.find_session(data, session_id)
    if not afc_enrolments.supported(services, training):
        abort(404)
    if payload.get("revision") != revision(services, training, data):
        return {"ok": False, "error": "La session a changé. Rechargez la page avant d’enregistrer.", "status": 409}
    try:
        dates = validate_dates(training, payload.get("dates"))
        rows = participants(services, training)
        choices = payload.get("profiles")
        if not isinstance(choices, dict) or set(choices) != {str(row["id"]) for row in rows}:
            raise ValueError("La liste des candidats a changé. Rechargez la page.")
        changes = []
        for trainee in rows:
            hours = choices[str(trainee["id"])]
            if hours is not None and (type(hours) is not int or hours not in PROFILES):
                raise ValueError("Choisissez 54 h, 40 h ou sans RAN pour chaque candidat.")
            values = {"afc_ran_hours": hours}
            if hours is not None:
                if not dates[str(hours)]:
                    raise ValueError(f"Renseignez la date d’entrée du profil à {hours} h.")
                values.update(afc_schedule.trainee_update(training, trainee, {
                    "training_date_start": dates[str(hours)]}))
            changes.append((trainee, values))
    except ValueError as exc:
        return {"ok": False, "error": str(exc), "status": 400}
    # Validate every row before writing anything. No documents or emails are sent.
    now = services._now_iso()
    proposed = suggestion(services, training, data)
    training["afc_hours"] = {"dates": dates, "updated_at": now,
                             "positioning_scores": proposed["scores"],
                             "allocation_method": "positioning_tests" if choices == proposed["profiles"] else "manual"}
    for trainee, values in changes:
        if values.get("training_date_start", trainee.get("training_date_start")) != trainee.get("training_date_start"):
            values["afc_dates_changed_at"] = now
        trainee.update(values)
    return {"ok": True, "state": view(services, training, data)}


def workbook(training, state):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    book = Workbook()
    sheet = book.active
    sheet.title = "Parcours AFC"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:G1")
    sheet["A1"] = "AFC · Répartition des heures de formation"
    sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color="153C37")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells("A2:G2")
    sheet["A2"] = str(training.get("name") or "Session AFC")
    sheet["A2"].data_type = "s"
    sheet["A2"].font = Font(name="Calibri", size=12, color="53645F")
    sheet.row_dimensions[2].height = 24
    sheet.merge_cells("A3:G3")
    sheet["A3"] = (f"{len(state['rows'])} candidats · "
                   f"{state['counts']['54']} à 54 h de RAN · "
                   f"{state['counts']['40']} à 40 h · {state['counts']['0']} sans RAN")
    sheet["A3"].font = Font(name="Calibri", size=11, color="53645F")
    headers = ["Nom", "Prénom", "Formation professionnelle (h)", "RAN – Remise à niveau (h)",
               "Préparation à l’après-formation (h)", "Soutien personnalisé (h)", "Date de début de formation"]
    for col, label in enumerate(headers, 1):
        cell = sheet.cell(5, col, label)
        cell.fill = PatternFill("solid", fgColor="153C37")
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[5].height = 46
    for index, row in enumerate(state["rows"], 6):
        values = [row["last_name"], row["first_name"], COMMON_HOURS[0], row["ran_hours"],
                  COMMON_HOURS[1], COMMON_HOURS[2], dt.date.fromisoformat(row["start"])]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(index, col, value)
            if col <= 2:
                cell.data_type = "s"  # Names must never become Excel formulas.
            cell.font = Font(name="Calibri", size=11, color="20372F", bold=col == 4)
            cell.alignment = Alignment(vertical="center", horizontal="left" if col <= 2 else "center")
            cell.fill = PatternFill("solid", fgColor="F1F6F4" if index % 2 == 0 else "FFFFFF")
            if 3 <= col <= 6:
                cell.number_format = '0" h"'
            if col == 7:
                cell.number_format = "dd/mm/yyyy"
        sheet.row_dimensions[index].height = 26
    last = 5 + len(state["rows"])
    table = Table(displayName="ParcoursAFC", ref=f"A5:G{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    sheet.add_table(table)
    for col, width in zip("ABCDEFG", (24, 22, 23, 24, 29, 23, 26)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "C6"
    sheet.print_title_rows = "1:5"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:G{last}"
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


def register(app, services):
    def require_training(session_id):
        data = services.load_data(run_background_tasks=False)
        training = services.find_session(data, session_id)
        if not afc_enrolments.supported(services, training):
            abort(404)
        return data, training

    @app.get("/admin/sessions/<session_id>/afc-hours")
    @services.admin_login_required
    def admin_afc_hours(session_id):
        data, training = require_training(session_id)
        response = app.make_response(render_template("admin_afc_hours.html", training=training,
            afc_state=view(services, training, data)))
        response.headers["Cache-Control"] = "private, no-store"
        return response

    @app.post("/api/admin/sessions/<session_id>/afc-hours")
    @services.admin_login_required
    @services.admin_write_required
    def api_afc_hours(session_id):
        payload = request.get_json(silent=True)
        if not isinstance(payload, dict):
            abort(400)
        token = session.get("afc_hours_csrf", "")
        if not token or not hmac.compare_digest(str(payload.get("csrf") or ""), token):
            abort(403)
        result = services._atomic_update_data(lambda data: update(services, data, session_id, payload))
        status = result.pop("status", 200)
        response = jsonify(result)
        response.headers["Cache-Control"] = "private, no-store"
        return response, status

    @app.get("/admin/sessions/<session_id>/afc-hours.xlsx")
    @services.admin_login_required
    def admin_afc_hours_excel(session_id):
        data, training = require_training(session_id)
        state = view(services, training, data)
        if request.args.get("revision") and request.args["revision"] != state["revision"]:
            return jsonify(ok=False, error="La session a changé. Rechargez la page avant l’export."), 409
        try:
            validate_dates(training, state["dates"])
            error = export_error(state)
            if error:
                raise ValueError(error)
        except ValueError as exc:
            return jsonify(ok=False, error=str(exc)), 422
        filename = secure_filename(f"AFC_heures_{training.get('name') or session_id}") + ".xlsx"
        response = send_file(workbook(training, state), as_attachment=True, download_name=filename,
                             mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response.headers["Cache-Control"] = "private, no-store"
        return response
