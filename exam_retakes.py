"""Independent exam retakes: never insert a second administrative enrolment.

SQLite checkpoints and per-dossier process locks protect remote side effects.
No provider call is made while rendering a trainee page.
"""
from __future__ import annotations

import contextlib
import datetime as dt
import fcntl
import hashlib
import hmac
import json
import os
import re
import secrets
import sqlite3
import tempfile
import threading
import time
import unicodedata
import uuid
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from zoneinfo import ZoneInfo

from flask import abort, jsonify, request, send_file, session, url_for


class RetakeError(ValueError):
    pass


def now():
    return dt.datetime.now(dt.timezone.utc).isoformat()


def today():
    return dt.datetime.now(ZoneInfo("Europe/Paris")).date().isoformat()


def folded(value):
    return "".join(c for c in unicodedata.normalize("NFKD", str(value or "")) if not unicodedata.combining(c)).lower().replace("’", "'")


def cents(value, label="Montant"):
    try:
        amount = Decimal(str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", "."))
        if not amount.is_finite() or amount < 0 or amount > 100000:
            raise InvalidOperation
        return int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError, TypeError):
        raise RetakeError(label + " invalide : indiquez un montant positif ou nul.") from None


def date_value(value, label):
    try:
        value = str(value or "")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError
        return dt.date.fromisoformat(value).isoformat()
    except ValueError:
        raise RetakeError(label + " : date invalide.") from None


def money(value):
    return f"{Decimal(value) / 100:.2f}".replace(".", ",") + " €"


def date_label(value):
    try:
        return dt.date.fromisoformat(value).strftime("%d/%m/%Y")
    except (ValueError, TypeError):
        return str(value or "À préciser")


def family(training):
    value = folded(training.get("training_type") or training.get("type_formation") or "").upper()
    if "VAE" in value:
        return value
    if "DESP" in value or "DIRIGEANT" in value:
        return "DESP"
    if "VTC" in value:
        return "VTC"
    match = re.search(r"SSIAP\s*([123])", value)
    if match:
        return "SSIAP " + match[1]
    return value.strip()


def components(host, training):
    kind = family(training)
    if host.afc_schedule.enabled(training):
        return {"APS": "Examen APS", "SSIAP 1": "Examen SSIAP 1"}
    if "VTC" in kind:
        return {"VTC pratique": "VTC · épreuve pratique", "VTC théorie": "VTC · épreuve théorique"}
    return {kind: "Examen " + kind}


def exam_date(host, training, kind):
    key = ("exam_practice_date" if kind == "VTC pratique" else "exam_theory_date" if kind == "VTC théorie"
           else "ssiap_exam_date" if kind.startswith("SSIAP") else "exam_date")
    # Never propose the APS date for the SSIAP component of an AFC session.
    value = host._session_get(training, key, "")
    if not value and not host.afc_schedule.enabled(training) and kind != "VTC pratique":
        value = host._session_get(training, "exam_date", "")
    try:
        return date_value(str(value or "")[:10], "Examen")
    except RetakeError:
        return ""


def preparation_days(host, training, kind):
    """Suggest the last two course weekdays of the receiving session, before its exam."""
    def parse(value):
        try:
            return dt.date.fromisoformat(date_value(value, "Formation"))
        except RetakeError:
            return None
    phase = "aps_in_person" if kind == "APS" and not host.afc_schedule.enabled(training) else "dirigeant_in_person" if kind == "DESP" else ""
    end = parse(training.get(phase + "_end")) if phase else None
    end = end or parse(host._session_get(training, "date_end", ""))
    exam = parse(exam_date(host, training, kind))
    if exam:
        end = min(end or exam, exam - dt.timedelta(days=1))
    if not end:
        return []
    start = parse(training.get(phase + "_start")) if phase else None
    start = start or parse(host._session_get(training, "date_start", ""))
    start = max(start or dt.date.fromisoformat(today()), dt.date.fromisoformat(today()))
    exams = {parse(exam_date(host, training, item)) for item in components(host, training)}
    dates = []
    for offset in range(14):
        day = end - dt.timedelta(days=offset)
        if day < start:
            break
        if day not in exams and host._is_aps_working_day(day):
            dates.append({"date": day.isoformat()})
            if len(dates) == 2:
                break
    return list(reversed(dates))


def candidates(host, data, source, kind):
    result = []
    cutoff = max(today(), exam_date(host, source, kind))
    tenant = source.get("partner_id") or host.INTEGRALE_PARTNER_ID
    for target in data.get("sessions", []):
        if target.get("archived") or target.get("id") == source.get("id") or host._is_wedof_leads_session(target):
            continue
        if (target.get("partner_id") or host.INTEGRALE_PARTNER_ID) != tenant or kind not in components(host, target):
            continue
        value = exam_date(host, target, kind)
        if value and value < cutoff:
            continue
        if not value and str(host._session_get(target, "date_end", "") or "") < today():
            continue
        result.append({"id": target["id"], "name": host._session_get(target, "name", "") or kind,
                       "exam_date": value, "date_label": date_label(value),
                       "preparation_days": preparation_days(host, target, kind)})
    return sorted(result, key=lambda row: (row["exam_date"] or "9999", row["name"]))


def read_contract(path):
    if path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(path)
        return "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
    return ""


def prices(host, training, trainee, kind):
    signature = host._yousign_state(trainee)
    sources = [(signature.get("signed_pdf_path"), "Convention signée du stagiaire"),
               (trainee.get("convention_aps_docx_path"), "Convention du stagiaire"),
               (signature.get("unsigned_pdf_path") or trainee.get("convention_aps_pdf_path"), "Convention du stagiaire")]
    config = host._automation_document_config(training)
    if config.get("convention_template") and (family(training) in {"APS", "A3P", "DESP", "SSIAP 1", "VTC", "DIRIGEANT VAE"} or host.afc_schedule.enabled(training)):
        sources.append((str(Path(host.app.root_path) / "templates_word" / config["convention_template"]), "Modèle de convention de la formation"))
    text, source = "", "Aucun tarif contractuel disponible"
    for raw, label in sources:
        if not raw or not Path(raw).is_file():
            continue
        try:
            text = read_contract(Path(raw))
            source = label
            # Do not silently substitute today's template for an unreadable signed contract.
            break
        except Exception:
            source = label + " : lecture impossible, tarif à renseigner"
            break
    normal = re.sub(r"\s+", " ", folded(text))
    exam = re.search(r"representation a l'examen est facturee\s*:\s*([\d ,\.]+)\s*€", normal)
    day = re.search(r"journee de rattrapage est facturee\s*:\s*([\d ,\.]+)\s*€", normal)
    note = ""
    if kind == "VTC pratique":
        exam = re.search(r"vehicule a double commande.{0,180}cout forfaitaire de\s*([\d ,\.]+)\s*€", normal)
        note = "Le forfait porte sur le véhicule à double commande. Les frais CMAR restent distincts et à la charge du stagiaire."
    elif kind == "VTC théorie":
        exam = None
        note = "Les frais de représentation CMAR ne sont pas inclus : renseignez uniquement la prestation facturée par Intégrale Academy."
    if host.afc_schedule.enabled(training):
        note = "Convention France Travail : aucun tarif personnel n’est présumé. Vérifiez les modalités convenues avant de facturer."
    return {"exam_price": f"{Decimal(cents(exam[1])) / 100:.2f}" if exam else "",
            "day_price": f"{Decimal(cents(day[1])) / 100:.2f}" if day else "", "source": source, "note": note,
            "exam_clause": exam[0] if exam else "Tarif d’examen non précisé dans la convention.",
            "day_clause": day[0] if day else "Tarif journalier non précisé dans la convention."}


class Store:
    def __init__(self, host):
        self.root = Path(host.PERSIST_DIR) / "exam_retakes"
        self.root.mkdir(parents=True, exist_ok=True, mode=0o700)
        self.db = self.root / "retakes.sqlite3"
        with self.connect() as db:
            db.executescript("""CREATE TABLE IF NOT EXISTS retakes (
                id TEXT PRIMARY KEY, source_session TEXT NOT NULL, trainee TEXT NOT NULL,
                target_session TEXT NOT NULL, status TEXT NOT NULL, request_id TEXT,
                invoice_id TEXT, next_attempt REAL NOT NULL, payload TEXT NOT NULL);
                CREATE INDEX IF NOT EXISTS retakes_trainee ON retakes(source_session,trainee);
                CREATE INDEX IF NOT EXISTS retakes_target ON retakes(target_session,status);
                CREATE UNIQUE INDEX IF NOT EXISTS retakes_request ON retakes(request_id) WHERE request_id IS NOT NULL;
                CREATE UNIQUE INDEX IF NOT EXISTS retakes_invoice ON retakes(invoice_id) WHERE invoice_id IS NOT NULL;
            """)

    def connect(self):
        return sqlite3.connect(self.db, timeout=15)

    @contextlib.contextmanager
    def lock(self, key):
        with (self.root / (hashlib.sha256(key.encode()).hexdigest() + ".lock")).open("a") as handle:
            try:
                fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
            except BlockingIOError:
                raise RetakeError("Une opération est en cours. Patientez puis actualisez le suivi.") from None
            yield

    def save(self, row):
        row["updated_at"] = now()
        with self.connect() as db:
            db.execute("INSERT OR REPLACE INTO retakes VALUES (?,?,?,?,?,?,?,?,?)", (
                row["id"], row["source_session"], row["trainee_id"], row.get("target_session", ""), row["status"],
                row.get("request_id") or None, row.get("invoice_id") or None, row.get("next_attempt", 0), json.dumps(row)))

    def get(self, rid):
        with self.connect() as db:
            row = db.execute("SELECT payload FROM retakes WHERE id=?", (rid,)).fetchone()
        return json.loads(row[0]) if row else None

    def query(self, where="1", args=()):
        with self.connect() as db:
            rows = db.execute("SELECT payload FROM retakes WHERE " + where + " ORDER BY rowid DESC", args).fetchall()
        return [json.loads(row[0]) for row in rows]

    def document(self, rid, kind):
        if not re.fullmatch(r"[a-f0-9]{32}", rid) or kind not in {"convention", "signed", "invoice"}:
            raise RetakeError("Document introuvable.")
        return self.root / (rid + "-" + kind + ".pdf")


def line_items(row):
    result = []
    exam_title = "Véhicule double commande · rattrapage VTC" if row["kind"] == "VTC pratique" else "Rattrapage examen · " + row["kind"]
    result.append({"label": exam_title + " · " + date_label(row["exam_date"]), "quantity": 1,
                   "unit_cents": row["exam_cents"], "offered": row["exam_offered"]})
    if row["days"]:
        result.append({"label": "Journées de préparation · " + row["kind"], "quantity": len(row["days"]),
                       "unit_cents": row["day_cents"], "offered": row["days_offered"]})
    reference = row.get("price_reference") or {}
    for item, key in zip(result, ("exam_price", "day_price")):
        try:
            base = cents(reference.get(key))
        except RetakeError:
            base = None
        unit, quantity = item["unit_cents"], item["quantity"]
        gross = max(unit, base or 0) * quantity
        item["total_cents"] = 0 if item["offered"] else item["unit_cents"] * item["quantity"]
        item["reference_unit_cents"] = base
        item["gross_cents"] = gross
        item["discount_cents"] = 0 if item["offered"] else max((base or 0) - unit, 0) * quantity
        item["offered_cents"] = gross if item["offered"] else 0
        item["discount_percent"] = str((Decimal(item['discount_cents']) * 100 / (base * quantity)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)) if item["discount_cents"] else "0.00"
    return result


def financial_totals(row):
    items = line_items(row)
    return {key: sum(item[key] for item in items) for key in ("gross_cents", "discount_cents", "offered_cents", "total_cents")}


def discount_label(item):
    if not item["discount_cents"]:
        return ""
    percent = item["discount_percent"].rstrip("0").rstrip(".").replace(".", ",")
    return "Remise de " + money(item["discount_cents"]) + " (" + percent + " %)"


def validate(host, data, source, trainee, payload):
    kind = str(payload.get("kind") or "")
    if kind not in components(host, source):
        raise RetakeError("Choisissez l’examen correspondant à la formation.")
    exam = date_value(payload.get("exam_date"), "Examen")
    if exam < today():
        raise RetakeError("La nouvelle date d’examen doit être à venir.")
    failed = date_value(payload.get("failed_exam_date") or exam_date(host, source, kind), "Examen non réussi")
    if failed > today() or failed >= exam:
        raise RetakeError("La date de l’examen non réussi doit être passée et antérieure au nouvel examen.")
    target = host.find_session(data, str(payload.get("target_session") or ""))
    if not target or target.get("archived") or kind not in components(host, target):
        raise RetakeError("Choisissez une session d’accueil de la même formation.")
    if (target.get("partner_id") or host.INTEGRALE_PARTNER_ID) != (source.get("partner_id") or host.INTEGRALE_PARTNER_ID):
        raise RetakeError("La session d’accueil appartient à un autre organisme.")
    days = payload.get("days") or []
    if not isinstance(days, list) or len(days) > 60:
        raise RetakeError("Les journées de préparation sont invalides (60 maximum).")
    no_days = payload.get("no_days") is True
    if no_days:
        days = []
    elif not days:
        raise RetakeError("Ajoutez une journée ou cochez « Pas de journées de cours supplémentaires ».")
    cleaned, seen = [], set()
    for item in days:
        if not isinstance(item, dict):
            raise RetakeError("Journée de préparation invalide.")
        value = date_value(item.get("date"), "Journée de préparation")
        if value < today() or value > exam or value in seen:
            raise RetakeError("Les journées doivent être distinctes, à venir et antérieures ou égales à l’examen.")
        start, end = str(item.get("start") or "09:00"), str(item.get("end") or "17:00")
        if not all(re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", x) for x in (start, end)) or start >= end:
            raise RetakeError("Vérifiez les horaires des journées de préparation.")
        seen.add(value)
        cleaned.append({"date": value, "start": start, "end": end, "topic": str(item.get("topic") or "Révision et préparation à l’examen").strip()[:300]})
    exam_offered, days_offered = payload.get("exam_offered") is True, payload.get("days_offered") is True
    result = {"kind": kind, "exam_date": exam, "failed_exam_date": failed,
              "failed_exam_details": str(payload.get("failed_exam_details") or "").strip()[:1000],
              "source_date_start": host._session_get(source, "date_start", ""),
              "source_date_end": host._session_get(source, "date_end", ""), "target_session": target["id"],
              "target_name": host._session_get(target, "name", "") or kind,
              "exam_time": str(payload.get("exam_time") or "").strip(),
              "location": str(payload.get("location") or "").strip()[:500],
              "notes": str(payload.get("notes") or "").strip()[:2000],
              "days": sorted(cleaned, key=lambda x: x["date"]), "no_days": no_days,
              "exam_cents": cents(payload.get("exam_price") or "0") if exam_offered else cents(payload.get("exam_price"), "Tarif de l’examen"),
              "day_cents": cents(payload.get("day_price") or "0") if days_offered or no_days else cents(payload.get("day_price"), "Tarif journalier"),
              "exam_offered": exam_offered, "days_offered": days_offered,
              "price_reference": prices(host, source, trainee, kind)}
    if not result["location"]:
        raise RetakeError("Précisez le lieu de l’examen et des journées de préparation.")
    if result["exam_time"] and not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", result["exam_time"]):
        raise RetakeError("Horaire d’examen invalide.")
    result["total_cents"] = sum(item["total_cents"] for item in line_items(result))
    return result


LABELS = {"draft": "Brouillon", "queued": "Préparation de la signature", "signing": "En attente de signature",
          "signed": "Convention signée · préparation de la facture", "payment": "En attente de paiement",
          "confirmed": "Rattrapage confirmé", "stopped": "Signature arrêtée", "cancelled": "Brouillon annulé"}


def view(row, *, public=False, token=""):
    totals = financial_totals(row)
    result = {k: row.get(k) for k in ("id", "status", "kind", "source_session", "trainee_id", "target_session", "target_name",
              "exam_date", "exam_time", "location", "days", "no_days", "notes", "exam_offered", "days_offered", "revision",
              "signed_at", "confirmed_at", "invoice_number", "payment_status", "created_at", "failed_exam_date", "failed_exam_details")}
    result.update(label=LABELS.get(row["status"], row["status"]), total=money(row["total_cents"]),
                  total_cents=row["total_cents"], exam_price=f"{row['exam_cents']/100:.2f}", day_price=f"{row['day_cents']/100:.2f}",
                  exam_date_label=date_label(row["exam_date"]), failed_exam_date_label=date_label(row.get("failed_exam_date")),
                  items=[{**x, "price": money(x["total_cents"]), "discount_label": discount_label(x),
                          "reference_price": money(x["reference_unit_cents"]) if x["reference_unit_cents"] is not None else ""} for x in line_items(row)],
                  financials={key: money(value) for key, value in totals.items()},
                  has_discount=totals["discount_cents"] > 0,
                  has_offered=totals["offered_cents"] > 0)
    result["signature_url"] = row.get("signature_url", "") if row["status"] == "signing" else ""
    result["payment_url"] = row.get("payment_url", "") if row["status"] == "payment" and row.get("payment_status") == "unpaid" else ""
    endpoint = "retake_public_document" if public else "retake_document"
    extra = {"token": token} if public else {}
    result["convention_url"] = url_for(endpoint, rid=row["id"], kind="signed" if row.get("signed_at") else "convention",
                                       v=row.get("pdf_sha256") or row.get("revision", 1), **extra)
    result["invoice_url"] = url_for(endpoint, rid=row["id"], kind="invoice", **extra) if row.get("invoice_pdf") else ""
    if not public:
        result.update(error=row.get("error", ""), price_reference=row.get("price_reference", {}),
                      needs_review=row["status"] == "draft" and row.get("document_version", 1) < 2,
                      email_sent=bool(row.get("emails", {}).get("documents", {}).get("sent_at")),
                      trainee_name=row["person"]["first_name"] + " " + row["person"]["last_name"])
    return result


def views(host, sid, tid, *, public=False, token=""):
    return [view(r, public=public, token=token) for r in Store(host).query("source_session=? AND trainee=?", (sid, tid))
            if not public or r["status"] not in {"draft", "cancelled", "queued"}]


def roster(host, sid):
    return [view(r) for r in Store(host).query("target_session=? AND status='confirmed'", (sid,))]


def scoped(host, sid, tid):
    data = host.load_data(run_background_tasks=False)
    source = host.find_session(data, sid)
    trainee = next((t for t in host._session_trainees_list(source) if t.get("id") == tid), None) if source else None
    if not source or not trainee:
        abort(404)
    if (source.get("partner_id") or host.INTEGRALE_PARTNER_ID) != host.INTEGRALE_PARTNER_ID:
        abort(403)
    return data, source, trainee


def preflight(host, row):
    person = row["person"]
    for field in ("first_name", "last_name", "email", "phone", "address", "zip_code", "city"):
        if not str(person.get(field) or "").strip():
            raise RetakeError("Complétez d’abord les coordonnées du stagiaire dans sa fiche (nom, prénom, e-mail, mobile et adresse).")
    if not re.fullmatch(r"[^\s<>@]+@[^\s<>@]+\.[^\s<>@]+", person["email"]):
        raise RetakeError("L’adresse e-mail du stagiaire est invalide.")
    if not host._normalize_yousign_sms_phone(person["phone"]):
        raise RetakeError("Un numéro de mobile valide est nécessaire pour la signature par SMS.")
    if not host._yousign_is_configured() or host._missing_brevo_config():
        raise RetakeError("Vérifiez les connexions Yousign et e-mail dans les réglages.")
    if row["total_cents"]:
        if not host._qonto_is_configured():
            raise RetakeError("La connexion Qonto doit être configurée avant l’envoi.")
        host.get_qonto_invoice_iban()
        data = host.load_data(run_background_tasks=False)
        if not host._qonto_oauth_connected(data) or not all(host._qonto_oauth_has_scope(x, data) for x in ("client_invoices.read", "payment_link.read", "payment_link.write")):
            raise RetakeError("Autorisez les liens de paiement dans les réglages Qonto avant l’envoi.")


def register(app, host):
    @app.get('/api/sessions/<sid>/stagiaires/<tid>/rattrapages')
    @host.admin_login_required
    def retake_list(sid, tid):
        data, source, trainee = scoped(host, sid, tid)
        if request.args.get('status') == '1':
            return jsonify(ok=True, records=views(host, sid, tid))
        defaults = []
        for kind, label in components(host, source).items():
            defaults.append({"kind": kind, "label": label, "prices": prices(host, source, trainee, kind),
                             "failed_exam_date": exam_date(host, source, kind) if exam_date(host, source, kind) <= today() else "",
                             "sessions": candidates(host, data, source, kind)})
        return jsonify(ok=True, records=views(host, sid, tid), defaults=defaults,
                       csrf=session.setdefault("retake_csrf", secrets.token_urlsafe(32)),
                       location="Intégrale Academy 54 chemin du Carreou 83480 PUGET SUR ARGENS",
                       read_only=session.get("admin_role") == "viewer")

    @app.post('/api/sessions/<sid>/stagiaires/<tid>/rattrapages')
    @host.admin_login_required
    @host.admin_write_required
    def retake_save(sid, tid):
        data, source, trainee = scoped(host, sid, tid)
        payload = request.get_json(silent=True) or {}
        check_csrf(payload)
        store = Store(host)
        try:
            with store.lock(sid + ":" + tid):
                rid = str(payload.get("id") or "")
                previous = store.get(rid) if rid else None
                if rid and (not previous or previous["source_session"] != sid or previous["trainee_id"] != tid):
                    abort(404)
                with store.lock(rid or uuid.uuid4().hex):
                    if rid:
                        previous = store.get(rid)
                    if previous and (previous["status"] != "draft" or payload.get("revision") != previous["revision"]):
                        raise RetakeError("Le dossier a évolué. Actualisez le suivi avant de le modifier.")
                    values_payload = payload
                    if payload.get("refresh_convention") is True:
                        if not previous:
                            raise RetakeError("Enregistrez d’abord le rattrapage avant d’actualiser sa convention.")
                        # Refresh the trainee/session details without replacing negotiated terms
                        # with today's defaults or accepting hidden changes in this request.
                        values_payload = {k: previous.get(k) for k in (
                            "kind", "target_session", "exam_date", "exam_time", "location", "notes",
                            "failed_exam_date", "failed_exam_details", "days", "no_days", "exam_offered", "days_offered")}
                        values_payload.update(exam_price=f"{previous['exam_cents'] / 100:.2f}",
                                              day_price=f"{previous['day_cents'] / 100:.2f}")
                    values = validate(host, data, source, trainee, values_payload)
                    if previous and previous["kind"] == values["kind"] and previous.get("price_reference"):
                        values["price_reference"] = previous["price_reference"]
                    if not rid and any(r["status"] not in {"cancelled", "stopped"} and r["kind"] == values["kind"] and r["exam_date"] >= today()
                                       for r in store.query("source_session=? AND trainee=?", (sid, tid))):
                        raise RetakeError("Un rattrapage est déjà prévu pour cet examen. Ouvrez le dossier existant.")
                    row = {**(previous or {}), **values, "id": rid or uuid.uuid4().hex, "status": "draft",
                           "source_session": sid, "source_name": host._session_get(source, "name", ""), "trainee_id": tid,
                           "person": {k: str(trainee.get(k) or "").strip() for k in ("first_name", "last_name", "email", "phone", "address", "zip_code", "city", "birth_date")},
                           "vat_rate": str(trainee.get("qonto_vat_rate") if trainee.get("qonto_vat_rate") not in (None, "") else "0"),
                           "created_at": (previous or {}).get("created_at", now()), "revision": (previous or {}).get("revision", 0) + 1,
                           "document_version": 2, "document_date": today(),
                           "actor": session.get("admin_username") or "Administration"}
                    from exam_retakes_documents import build_convention
                    destination = store.document(row["id"], "convention")
                    # Keep the previous PDF intact if generation fails midway through.
                    with tempfile.NamedTemporaryFile(dir=store.root, suffix=".pdf", delete=False) as handle:
                        staged = Path(handle.name)
                    try:
                        row["signature_field"] = build_convention(row, staged)
                        row["pdf_sha256"] = hashlib.sha256(staged.read_bytes()).hexdigest()
                        original = destination.read_bytes() if destination.is_file() else None
                        os.replace(staged, destination)
                        try:
                            store.save(row)
                        except Exception:
                            if original is None:
                                destination.unlink(missing_ok=True)
                            else:
                                staged.write_bytes(original)
                                os.replace(staged, destination)
                            raise
                    finally:
                        staged.unlink(missing_ok=True)
            return jsonify(ok=True, record=view(row))
        except (RetakeError, ValueError) as exc:
            return jsonify(ok=False, error=str(exc)), 400

    @app.post('/api/rattrapages/<rid>/<action>')
    @host.admin_login_required
    @host.admin_write_required
    def retake_action(rid, action):
        payload = request.get_json(silent=True) or {}
        check_csrf(payload)
        store = Store(host)
        row = store.get(rid)
        if not row:
            abort(404)
        data, source, trainee = scoped(host, row["source_session"], row["trainee_id"])
        try:
            with store.lock(rid):
                row = store.get(rid)
                if action == "send":
                    if row["status"] != "draft" or payload.get("revision") != row["revision"]:
                        raise RetakeError("Le dossier a évolué ou a déjà été envoyé. Actualisez le suivi.")
                    if row.get("document_version", 1) < 2:
                        raise RetakeError("Complétez la date de l’examen non réussi puis enregistrez la nouvelle convention avant l’envoi.")
                    if row["exam_date"] < today() or any(d["date"] < today() for d in row["days"]):
                        raise RetakeError("Actualisez les dates avant l’envoi.")
                    target = host.find_session(data, row['target_session'])
                    if (not target or target.get('archived') or row['kind'] not in components(host, target)
                            or (target.get('partner_id') or host.INTEGRALE_PARTNER_ID) != (source.get('partner_id') or host.INTEGRALE_PARTNER_ID)):
                        raise RetakeError("La session d’accueil a changé. Actualisez le brouillon avant l’envoi.")
                    # The reviewed PDF and the intended recipient must still match.
                    if any(str(trainee.get(k) or "").strip() != value for k, value in row["person"].items()):
                        raise RetakeError("Les coordonnées du stagiaire ont changé. Enregistrez à nouveau le brouillon.")
                    preflight(host, row)
                    row.update(status="queued", next_attempt=0, error="", submitted_at=now())
                elif action == "refresh":
                    row.update(next_attempt=0, error="")
                elif action == "cancel" and row["status"] == "draft":
                    row.update(status="cancelled")
                else:
                    raise RetakeError("Cette action n’est pas disponible pour ce dossier.")
                store.save(row)
            wake_worker(host)
            return jsonify(ok=True, record=view(row))
        except (RetakeError, ValueError) as exc:
            return jsonify(ok=False, error=str(exc)), 409

    @app.get('/admin/rattrapages/<rid>/<kind>.pdf')
    @host.admin_login_required
    def retake_document(rid, kind):
        store = Store(host)
        row = store.get(rid)
        if not row:
            abort(404)
        scoped(host, row["source_session"], row["trainee_id"])
        return document_response(store, row, kind)

    @app.get('/espace/<token>/rattrapages/<rid>/<kind>.pdf')
    def retake_public_document(token, rid, kind):
        data = host.load_data(run_background_tasks=False)
        source, trainee = host.find_session_and_trainee_by_token(data, token)
        if not source or not trainee or not host._public_is_authed(token):
            abort(403)
        store = Store(host)
        row = store.get(rid)
        if not row or row["source_session"] != source["id"] or row["trainee_id"] != trainee["id"] or row["status"] in {"draft", "queued", "cancelled"}:
            abort(404)
        return document_response(store, row, kind)


def check_csrf(payload):
    expected = str(session.get("retake_csrf") or "")
    if not expected or not hmac.compare_digest(expected, str(payload.get("csrf") or "")):
        abort(403)


def document_response(store, row, kind):
    if kind not in {"convention", "signed", "invoice"}:
        abort(404)
    path = store.document(row["id"], kind)
    if not path.is_file():
        abort(404)
    response = send_file(path, mimetype="application/pdf", download_name="rattrapage-" + kind + ".pdf", conditional=False)
    response.headers["Cache-Control"] = "private, no-store"
    return response


def notify(host, *, request_id="", invoice_id=""):
    if not request_id and not invoice_id:
        return False
    store = Store(host)
    rows = store.query("request_id=?" if request_id else "invoice_id=?", (request_id or invoice_id,))
    if not rows:
        return False
    # Do not hold a webhook request while calling either provider. A scheduled
    # read also reconciles callbacks that arrive before their IDs are saved.
    for row in rows:
        with store.connect() as db:
            db.execute("UPDATE retakes SET next_attempt=0 WHERE id=?", (row["id"],))
    wake_worker(host)
    return True


def process_pending(host):
    from exam_retakes_workflow import process
    store = Store(host)
    for row in store.query("status IN ('queued','signing','signed','payment','confirmed') AND next_attempt<=?", (time.time(),))[:20]:
        try:
            with store.lock(row["id"]):
                current = store.get(row["id"])
                try:
                    process(host, store, current)
                    current["error"] = ""
                except Exception as exc:
                    current["error"] = host._sanitize_qonto_error(host._sanitize_yousign_error(str(exc)))[:700]
                    host.app.logger.warning("retake_processing_failed id=%s type=%s", current["id"], type(exc).__name__)
                current["next_attempt"] = time.time() + (86400 if current["status"] == "confirmed" and not current.get("error") else 300)
                store.save(current)
        except RetakeError:
            continue


def start_worker(host):
    if host.app.testing or "PYTEST_CURRENT_TEST" in os.environ:
        return
    key = "exam_retakes_worker"
    if host.app.extensions.get(key):
        return
    stop, wake = threading.Event(), threading.Event()
    def run():
        while not stop.is_set():
            try:
                with host.app.app_context():
                    process_pending(host)
            except Exception as exc:
                host.app.logger.error("retake_worker_failed type=%s", type(exc).__name__)
            wake.wait(45)
            wake.clear()
    thread = threading.Thread(target=run, name="exam-retakes", daemon=True)
    host.app.extensions[key] = {"stop": stop, "wake": wake, "thread": thread}
    thread.start()


def wake_worker(host):
    start_worker(host)
    state = host.app.extensions.get("exam_retakes_worker")
    if state:
        state["wake"].set()
