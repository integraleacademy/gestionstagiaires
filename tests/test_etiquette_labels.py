import unittest
import html
import re
import zipfile
from io import BytesIO
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

import app as gestion_app


class EtiquetteLabelsTests(unittest.TestCase):
    def test_template_mapping_matches_requested_training_types(self):
        self.assertEqual(gestion_app._etiquette_template_name("APS"), "etiquette_aps.docx")
        self.assertEqual(gestion_app._etiquette_template_name("A3P"), "etiquette_a3p.docx")
        self.assertEqual(gestion_app._etiquette_template_name("DIRIGEANT"), "etiquette_dirigeant_initial.docx")
        self.assertEqual(gestion_app._etiquette_template_name("DIRIGEANT VAE"), "etiquette_dirigeant.docx")
        self.assertEqual(gestion_app._etiquette_template_name(" ssiap 1 "), "etiquette_ssiap.docx")

    def _ssiap_session(self, trainee):
        return {
            "id": "SS1", "name": "SSIAP 1 février 2027", "training_type": "SSIAP 1",
            "date_start": "2027-02-11", "date_end": "2027-02-26", "trainees": [trainee],
        }

    def _admin_client(self):
        client = gestion_app.app.test_client()
        with client.session_transaction() as sess:
            sess["admin_logged_in"] = True
        return client

    def _afc_session(self, trainees):
        return {
            "id": "AFC1", "name": "AFC APS + SSIAP novembre 2026", "training_type": "APS", "afc_enabled": True,
            "date_start": "2026-11-02", "date_end": "2027-01-15", "trainees": trainees,
        }

    def test_afc_print_and_word_share_individual_dates_and_dossier_statuses(self):
        trainee = {
            "id": "TAFC", "last_name": "Exemple", "first_name": "Camille",
            "training_date_start": "2026-11-09", "training_date_end": "2026-12-30",
            "convention_status": "signing", "convention_signature": {"status": "done"},
            "test_fr_status": "validated",
            "documents": [
                {"key": "id", "status": "CONFORME"},
                {"key": "carte_vitale_doc", "status": "A CONTRÔLER"},
                {"key": "photo", "status": "NON CONFORME"},
                {"key": "cnaps_doc", "status": "CONFORME"},
                {"key": "certificat_medical_ssiap_afc", "status": " conforme "},
                {"key": "certificat_medical_ssiap", "status": "NON CONFORME"},
                {"key": "certif_med", "status": "NON CONFORME"},
            ],
        }
        expected = [
            (True, "Convention de formation signée"), (True, "Pièce d’identité"),
            (False, "Carte vitale"), (False, "Photo d’identité"), (True, "Autorisation préalable CNAPS"),
            (False, "Fiche CNIL signée"), (True, "Test de français"), (False, "SST à jour"),
            (True, "Certificat médical de moins de 3 mois"),
            (False, "Dossier examen APS"), (False, "Dossier examen SSIAP 1"),
        ]
        data = {"sessions": [self._afc_session([trainee])]}
        client = self._admin_client()
        route = "/admin/sessions/AFC1/stagiaires/TAFC/etiquette"
        with patch.object(gestion_app, "load_data", return_value=data), patch.object(gestion_app, "save_data"):
            response = client.get(route)
            self.assertEqual(response.status_code, 200)
            page = response.get_data(as_text=True)
            rows = re.findall(r'<li><span class="checklist-prefix">(.*?)</span><span>(.*?)</span></li>', page)
            self.assertEqual([(mark == "✔️", html.unescape(label)) for mark, label in rows], expected)
            word_response = client.get(route + ".docx")
        self.assertEqual(word_response.status_code, 200)
        doc = Document(BytesIO(word_response.data))
        paragraphs = ["".join(node.xpath(".//w:t/text()")) for node in doc.element.xpath(".//w:p")]
        word_text = "\n".join(paragraphs)
        for text in (page, word_text):
            for content in ("FORMATION AFC", "APS + SSIAP 1 + SST", "Habilitation électrique", "09/11/2026 → 15/01/2027"):
                self.assertIn(content, text)
            self.assertNotIn("02/11/2026", text)
            self.assertNotIn("30/12/2026", text)
            self.assertNotIn("{{", text)
        for checked, label in expected:
            self.assertIn(f"{'✓' if checked else '○'}  {label}", paragraphs)
        self.assertEqual(trainee["training_date_start"], "2026-11-09")
        for section in doc.sections:
            source = section._sectPr.find(qn("w:paperSrc"))
            self.assertEqual(source.get(qn("w:first")), "261")
            self.assertEqual(source.get(qn("w:other")), "261")

    def test_afc_bulk_word_uses_each_start_and_the_common_end(self):
        trainees = [
            {"id": "T1", "last_name": "Exemple", "first_name": "Camille", "training_date_start": "2026-11-09"},
            {"id": "T2", "last_name": "Test", "first_name": "Alex", "training_date_end": "2026-11-10"},
        ]
        data = {"sessions": [self._afc_session(trainees)]}
        with patch.object(gestion_app, "load_data", return_value=data), patch.object(gestion_app, "save_data"):
            response = self._admin_client().get("/admin/sessions/AFC1/stagiaires/etiquettes-word.zip")
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(BytesIO(response.data)) as archive:
            self.assertEqual(len(archive.namelist()), 2)
            for filename in archive.namelist():
                doc = Document(BytesIO(archive.read(filename)))
                text = "\n".join(doc.element.xpath(".//w:t/text()"))
                start = "09/11/2026" if "T1" in filename else "02/11/2026"
                self.assertIn(start + " → 15/01/2027", text)
                self.assertIn("FORMATION AFC", text)
                self.assertEqual(text.count("○"), 11)
                self.assertNotIn("{{", text)

    def test_explicit_afc_opt_out_preserves_the_aps_label(self):
        trainee = {"id": "T1", "last_name": "Exemple", "first_name": "Camille", "training_date_start": "2026-11-09"}
        training = self._afc_session([trainee])
        training["afc_enabled"] = False
        text = "\n".join(Document(BytesIO(gestion_app._build_etiquette_docx_bytes(training, trainee))).element.xpath(".//w:t/text()"))
        self.assertIn("FORMATION TFP APS", text)
        self.assertIn("02/11/2026 → 15/01/2027", text)
        self.assertNotIn("Dossier examen SSIAP 1", text)
        self.assertNotIn("FORMATION AFC", text)

    def test_ssiap_print_and_word_show_the_same_document_statuses(self):
        trainee = {
            "id": "TSSIAP", "last_name": "Exemple", "first_name": "Camille",
            "convention_status": "signed",
            "documents": [
                {"key": "id", "status": "CONFORME"},
                {"key": "carte_vitale_doc", "status": "A CONTRÔLER", "file": "vitale.pdf"},
                {"key": "photo", "status": "NON CONFORME", "file": "photo.png"},
                {"key": "certificat_medical_ssiap", "status": " conforme "},
                {"key": "attestation_secourisme_ssiap", "status": "CONFORME"},
                {"key": "certif_med", "status": "NON CONFORME"},
                {"key": "cnaps_doc", "status": "CONFORME"},
            ],
        }
        expected = [
            (True, "Contrat de formation signé"), (True, "Pièce d’identité"),
            (False, "Carte vitale"), (False, "Photo d’identité"),
            (True, "Certificat médical de moins de 3 mois"),
            (True, "Attestation secourisme (PSC1 ou SST)"),
            (False, "Dossier examen"),
        ]
        data = {"sessions": [self._ssiap_session(trainee)]}
        client = self._admin_client()
        route = "/admin/sessions/SS1/stagiaires/TSSIAP/etiquette"
        with patch.object(gestion_app, "load_data", return_value=data), patch.object(gestion_app, "save_data"):
            response = client.get(route)
            self.assertEqual(response.status_code, 200)
            page = response.get_data(as_text=True)
            rows = re.findall(r'<li><span class="checklist-prefix">(.*?)</span><span>(.*?)</span></li>', page)
            self.assertEqual([(mark == "✔️", html.unescape(label)) for mark, label in rows], expected)
            self.assertIn("FORMATION SSIAP 1", page)
            self.assertNotIn("Autorisation préalable CNAPS", page)
            self.assertNotIn("Test de français", page)
            self.assertNotIn("CNIL", page)

            word_response = client.get(route + ".docx")
            self.assertEqual(word_response.status_code, 200)
            doc = Document(BytesIO(word_response.data))
            texts = ["".join(node.xpath(".//w:t/text()")) for node in doc.element.xpath(".//w:p")]
            for checked, label in expected:
                self.assertIn(f"{'✓' if checked else '○'}  {label}", texts)
            all_text = "\n".join(texts)
            self.assertIn("FORMATION SSIAP 1", all_text)
            self.assertIn("11/02/2027 → 26/02/2027", all_text)
            self.assertNotIn("APS", all_text)
            self.assertNotIn("CNAPS", all_text)
            self.assertNotIn("CNIL", all_text)
            self.assertNotIn("{{", all_text)
            for section in doc.sections:
                source = section._sectPr.find(qn("w:paperSrc"))
                self.assertEqual(source.get(qn("w:first")), "261")
                self.assertEqual(source.get(qn("w:other")), "261")

    def test_ssiap_empty_dossier_keeps_all_unchecked_items_visible(self):
        trainee = {"id": "TSSIAP", "last_name": "Exemple", "first_name": "Camille"}
        data = {"sessions": [self._ssiap_session(trainee)]}
        client = self._admin_client()
        with patch.object(gestion_app, "load_data", return_value=data), patch.object(gestion_app, "save_data"):
            response = client.get("/admin/sessions/SS1/stagiaires/TSSIAP/etiquette")
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertEqual(page.count('class="checklist-prefix">○</span>'), 7)
        self.assertIn("Certificat médical de moins de 3 mois", page)
        self.assertIn("Attestation secourisme (PSC1 ou SST)", page)

    def test_ssiap_bulk_word_export_includes_each_trainee(self):
        training = self._ssiap_session({"id": "T1", "last_name": "Exemple", "first_name": "Camille"})
        training["trainees"].append({"id": "T2", "last_name": "Exemple", "first_name": "Alex"})
        client = self._admin_client()
        with patch.object(gestion_app, "load_data", return_value={"sessions": [training]}), patch.object(gestion_app, "save_data"):
            response = client.get("/admin/sessions/SS1/stagiaires/etiquettes-word.zip")
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(BytesIO(response.data)) as archive:
            self.assertEqual(len(archive.namelist()), 2)
            for filename in archive.namelist():
                doc = Document(BytesIO(archive.read(filename)))
                text = "\n".join(doc.element.xpath(".//w:t/text()"))
                self.assertIn("FORMATION SSIAP 1", text)
                self.assertEqual(text.count("○"), 7)
                self.assertNotIn("CNIL", text)
                self.assertNotIn("{{", text)

    def test_print_page_displays_expected_checklist_for_a3p_and_vae(self):
        fake_data = {
            "sessions": [
                {
                    "id": "S1",
                    "name": "Formation test",
                    "training_type": "A3P",
                    "date_start": "2026-01-05",
                    "date_end": "2026-02-06",
                    "trainees": [{"id": "T1", "last_name": "Doe", "first_name": "Jane"}],
                },
                {
                    "id": "S2",
                    "name": "Formation test VAE",
                    "training_type": "DIRIGEANT VAE",
                    "date_start": "2026-01-05",
                    "date_end": "2026-02-06",
                    "trainees": [{"id": "T2", "last_name": "Doe", "first_name": "John"}],
                },
            ]
        }

        client = gestion_app.app.test_client()
        with client.session_transaction() as sess:
            sess["admin_logged_in"] = True

        with patch.object(gestion_app, "load_data", return_value=fake_data), patch.object(gestion_app, "save_data"):
            a3p_resp = client.get("/admin/sessions/S1/stagiaires/T1/etiquette")
            self.assertEqual(a3p_resp.status_code, 200)
            a3p_html = a3p_resp.get_data(as_text=True)
            self.assertIn("Permis de conduire valide", a3p_html)
            self.assertIn("Certificat médical de moins de 3 mois", a3p_html)

            vae_resp = client.get("/admin/sessions/S2/stagiaires/T2/etiquette")
            self.assertEqual(vae_resp.status_code, 200)
            vae_html = vae_resp.get_data(as_text=True)
            self.assertIn("Convention de VAE signée", vae_html)
            self.assertIn("Attestation de recevabilité", vae_html)
            self.assertIn("FORMATION DESP", vae_html)
            self.assertIn("FORMATION DESP VAE", vae_html)

    def test_print_page_uses_desp_label_for_dirigeant_initial_and_expected_highlight(self):
        fake_data = {
            "sessions": [
                {
                    "id": "S3",
                    "name": "Formation dirigeant",
                    "training_type": "DIRIGEANT initial",
                    "date_start": "2026-03-09",
                    "date_end": "2026-04-24",
                    "trainees": [{"id": "T3", "last_name": "Koita", "first_name": "Manoury"}],
                },
            ]
        }

        client = gestion_app.app.test_client()
        with client.session_transaction() as sess:
            sess["admin_logged_in"] = True

        with patch.object(gestion_app, "load_data", return_value=fake_data), patch.object(gestion_app, "save_data"):
            resp = client.get("/admin/sessions/S3/stagiaires/T3/etiquette")
            self.assertEqual(resp.status_code, 200)
            html = resp.get_data(as_text=True)
            self.assertIn("FORMATION DESP", html)
            self.assertNotIn("FORMATION DIRIGEANT INITIAL", html)
            self.assertIn("color:#fff", html)
            self.assertIn("background:#8b5e3c", html)



    def test_print_page_applies_training_specific_colors(self):
        fake_data = {
            "sessions": [
                {
                    "id": "S4",
                    "name": "Formation aps",
                    "training_type": "APS",
                    "date_start": "2026-04-01",
                    "date_end": "2026-06-26",
                    "trainees": [{"id": "T4", "last_name": "Oufqih", "first_name": "Yannis"}],
                },
                {
                    "id": "S5",
                    "name": "Formation a3p",
                    "training_type": "A3P",
                    "date_start": "2026-09-01",
                    "date_end": "2026-10-27",
                    "trainees": [{"id": "T5", "last_name": "Urbanik", "first_name": "Anthony"}],
                },
            ]
        }

        client = gestion_app.app.test_client()
        with client.session_transaction() as sess:
            sess["admin_logged_in"] = True

        with patch.object(gestion_app, "load_data", return_value=fake_data), patch.object(gestion_app, "save_data"):
            aps_resp = client.get("/admin/sessions/S4/stagiaires/T4/etiquette")
            self.assertEqual(aps_resp.status_code, 200)
            aps_html = aps_resp.get_data(as_text=True)
            self.assertIn("FORMATION TFP APS", aps_html)
            self.assertIn("background:#7dd3fc", aps_html)

            a3p_resp = client.get("/admin/sessions/S5/stagiaires/T5/etiquette")
            self.assertEqual(a3p_resp.status_code, 200)
            a3p_html = a3p_resp.get_data(as_text=True)
            self.assertIn("FORMATION A3P", a3p_html)
            self.assertIn("background:#fde047", a3p_html)
            self.assertIn("checklist-compact", a3p_html)

if __name__ == "__main__":
    unittest.main()
