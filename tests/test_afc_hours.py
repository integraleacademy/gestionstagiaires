import copy
import datetime as dt
import io
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

import app
import afc_hours
import afc_schedule


@pytest.fixture
def case(monkeypatch):
    training = {"id": afc_hours.PLANNING_SESSION, "name": "AFC APS + SSIAP NOVEMBRE 2026",
                "training_type": "APS", "date_start": "2026-11-16", "date_end": "2027-02-15",
                "trainees": [{"id": f"T{i:02}", "last_name": f"CANDIDAT{i:02}",
                              "first_name": "Émilie", "training_date_start": "2026-11-16"}
                             for i in range(12)]}
    data = {"sessions": [training]}
    monkeypatch.setattr(app, "load_data", lambda *a, **kw: copy.deepcopy(data))
    monkeypatch.setattr(app, "_atomic_update_data", lambda fn: fn(data))
    mail = Mock()
    monkeypatch.setattr(app, "brevo_send_email", mail)
    client = app.app.test_client()
    with client.session_transaction() as session:
        session["admin_logged_in"] = True
        session["admin_role"] = "admin"
    return client, training, data, mail


def urls(training):
    stem = f"/admin/sessions/{training['id']}/afc-hours"
    return stem, "/api" + stem, stem + ".xlsx"


def payload(client, training):
    client.get(urls(training)[0])
    with client.session_transaction() as session:
        csrf = session["afc_hours_csrf"]
    return {"csrf": csrf, "revision": afc_hours.revision(app, training),
            "dates": dict(afc_hours.PLANNING_DATES),
            "profiles": {f"T{i:02}": 54 if i < 8 else 40 if i < 10 else 0 for i in range(12)}}


def test_named_profiles_start_dates_and_xlsx(case):
    client, training, data, mail = case
    page, api, excel = urls(training)
    request = payload(client, training)
    before = client.get(excel)
    assert before.status_code == 422
    assert all("afc_ran_hours" not in row for row in training["trainees"])
    result = client.post(api, json=request)
    assert result.status_code == 200, result.json
    state = result.json["state"]
    assert state["counts"] == {"54": 8, "40": 2, "0": 2}
    assert not afc_hours.export_error(state)
    for row in training["trainees"]:
        assert row["training_date_start"] == request["dates"][str(row["afc_ran_hours"])]
        assert afc_schedule.effective_session(training, row)["date_end"] == "2027-02-15"
    response = client.get(excel + "?revision=" + state["revision"])
    assert response.status_code == 200
    assert response.headers["Content-Type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in response.headers["Content-Disposition"]
    book = load_workbook(io.BytesIO(response.data))
    sheet = book.active
    assert sheet.max_column == 7 and sheet.max_row == 17
    assert sheet.freeze_panes == "C6"
    assert sheet.tables["ParcoursAFC"].ref == "A5:G17"
    rows = list(sheet.iter_rows(min_row=6, values_only=True))
    assert all(row[2] == 273 and row[4] == 21 and row[5] == 45 for row in rows)
    assert [row[3] for row in rows] == [54] * 8 + [40] * 2 + [0] * 2
    assert [row[6] for row in rows] == ([dt.datetime(2026, 11, 16)] * 8
                                     + [dt.datetime(2026, 11, 18)] * 2 + [dt.datetime(2026, 11, 26)] * 2)
    assert sum(row[2] + row[3] + row[4] + row[5] for row in rows) == 4580
    assert all(sheet.cell(i, 4).data_type == "n" for i in range(6, 18))
    assert sheet["G6"].number_format == "dd/mm/yyyy"
    mail.assert_not_called()


@pytest.mark.parametrize("kind", ["unknown_id", "missing_id", "invalid_hours", "boolean", "bad_date", "out_of_session", "order"])
def test_invalid_input_changes_nothing(case, kind):
    client, training, data, mail = case
    request = payload(client, training)
    if kind == "unknown_id": request["profiles"]["UNKNOWN"] = 54
    if kind == "missing_id": request["profiles"].pop("T11")
    if kind == "invalid_hours": request["profiles"]["T11"] = 41
    if kind == "boolean": request["profiles"]["T11"] = False
    if kind == "bad_date": request["dates"]["0"] = "2026-11-31"
    if kind == "out_of_session": request["dates"]["0"] = "2027-03-01"
    if kind == "order": request["dates"]["40"] = "2026-11-27"
    before = copy.deepcopy(data)
    response = client.post(urls(training)[1], json=request)
    assert response.status_code == 400, response.json
    assert data == before
    mail.assert_not_called()


def test_drafts_allowed_but_incomplete_or_wrong_counts_cannot_export(case):
    client, training, data, mail = case
    request = payload(client, training)
    request["profiles"]["T11"] = None
    assert client.post(urls(training)[1], json=request).status_code == 200
    assert client.get(urls(training)[2]).status_code == 422
    request = payload(client, training)
    request["profiles"]["T11"] = 54
    assert client.post(urls(training)[1], json=request).status_code == 200
    response = client.get(urls(training)[2])
    assert response.status_code == 422
    assert "8 candidats" in response.json["error"]


@pytest.mark.parametrize("change", ["profile", "date", "roster", "name", "session_date"])
def test_stale_save_rejected(case, change):
    client, training, data, mail = case
    request = payload(client, training)
    if change == "profile": training["trainees"][0]["afc_ran_hours"] = 40
    if change == "date": training["trainees"][0]["training_date_start"] = "2026-11-23"
    if change == "roster": training["trainees"].pop()
    if change == "name": training["trainees"][0]["first_name"] = "Alex"
    if change == "session_date": training["date_start"] = "2026-11-15"
    before = copy.deepcopy(data)
    assert client.post(urls(training)[1], json=request).status_code == 409
    assert data == before


def test_dates_changed_elsewhere_block_stale_export(case):
    client, training, data, mail = case
    result = client.post(urls(training)[1], json=payload(client, training))
    old_revision = result.json["state"]["revision"]
    training["trainees"][0]["training_date_start"] = "2026-11-23"
    assert client.get(urls(training)[2] + "?revision=" + old_revision).status_code == 409
    response = client.get(urls(training)[2])
    assert response.status_code == 422 and "date individuelle" in response.json["error"]


def test_access_control_and_csrf(case):
    client, training, data, mail = case
    request = payload(client, training)
    wrong = {**request, "csrf": "incorrect"}
    assert client.post(urls(training)[1], json=wrong).status_code == 403
    with client.session_transaction() as session:
        session["admin_role"] = "viewer"
    assert client.post(urls(training)[1], json=request).status_code == 403
    assert client.get(urls(training)[0]).status_code == 200
    with client.session_transaction() as session:
        session.clear()
    assert client.post(urls(training)[1], json=request).status_code in (401, 302)
    assert client.get(urls(training)[2]).status_code in (401, 302)


@pytest.mark.parametrize("other", ["aps", "partner"])
def test_afc_only_and_integrale_scope(case, other):
    client, training, data, mail = case
    request = payload(client, training)
    if other == "aps": training["afc_enabled"] = False
    else: training["partner_id"] = "OTHER"
    assert client.get(urls(training)[0]).status_code == 404
    assert client.get(urls(training)[2]).status_code == 404
    assert client.post(urls(training)[1], json=request).status_code == 404
    page = client.get(f"/admin/sessions/{training['id']}/trainees")
    assert b"btnAfcHoursExport" not in page.data


def test_other_session_has_no_november_dates_or_quotas(case):
    client, training, data, mail = case
    training["id"] = "AFC-OTHER"
    training["date_start"] = "2027-03-01"
    settings = afc_hours.config(training)
    assert settings == {"dates": {"54": "2027-03-01", "40": "", "0": ""}, "target": None}


def test_excel_names_are_text_not_formulas(case):
    client, training, data, mail = case
    training["trainees"][0]["last_name"] = '=HYPERLINK("https://example.test")'
    assert client.post(urls(training)[1], json=payload(client, training)).status_code == 200
    response = client.get(urls(training)[2])
    book = load_workbook(io.BytesIO(response.data))
    formula_name = next(row[0] for row in book.active.iter_rows(min_row=6) if str(row[0].value).startswith("="))
    assert formula_name.data_type == "s"


def test_cancelled_registrations_excluded_and_profiles_survive_reload(case):
    client, training, data, mail = case
    training["trainees"].append({"id": "CANCELLED", "last_name": "ANNULÉ", "registration_cancelled": True})
    response = client.post(urls(training)[1], json=payload(client, training))
    assert response.status_code == 200, response.json
    assert len(response.json["state"]["rows"]) == 12
    html = client.get(urls(training)[0]).get_data(as_text=True)
    assert '"ran_hours": 40' in html and '"ran_hours": 0' in html
    assert 'data-trainee="CANCELLED"' not in html


def test_test_based_allocation_is_previewed_then_saved_and_new_scores_invalidate_form(case):
    client, training, data, mail = case
    points = [44, 39, 38, 39, 37, 42, 38, 41, 48, 49, 49, 45]
    data["positioning_tests"] = [{"id": f"TEST{i}", "contact": dict(row),
                                  "score": points[i], "total": 60,
                                  "created_at": "2026-10-14T10:00:00Z"}
                                 for i, row in enumerate(training["trainees"])]
    request = payload(client, training)
    request["revision"] = afc_hours.revision(app, training, data)
    proposed = afc_hours.suggestion(app, training, data)
    assert proposed["profiles"]["T11"] == 40 and proposed["profiles"]["T10"] == 0
    html = client.get(urls(training)[0]).get_data(as_text=True)
    assert '49/60' in html and 'afcHoursSuggest' in html
    assert all("afc_ran_hours" not in row for row in training["trainees"])
    request["profiles"] = proposed["profiles"]
    data["positioning_tests"][0]["score"] = 43
    assert client.post(urls(training)[1], json=request).status_code == 409
    request["revision"] = afc_hours.revision(app, training, data)
    response = client.post(urls(training)[1], json=request)
    assert response.status_code == 200, response.json
    assert training["afc_hours"]["allocation_method"] == "positioning_tests"
    assert training["afc_hours"]["positioning_scores"]["T00"]["score"] == 43
    assert client.get(urls(training)[2]).status_code == 200
    mail.assert_not_called()
