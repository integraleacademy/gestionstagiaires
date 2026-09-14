"""Official-reference examples, financial edge cases and protected form saves."""
import copy
import datetime as dt
import io
import tempfile
import unittest
from pathlib import Path

from docx import Document
from bts_cerfa import source_version, effective_values
from bts_contract_documents import defaults, validate_settings, fingerprint, fill_convention, document_errors
from bts_contract_routes import contract_context
from bts_contract_store import ContractStore
from bts_npec import quote, periods, reference
from bts_workspace import register_bts_workspace
from bts_workspace_store import WorkspaceError
from tests.test_bts_cerfa import complete_values
from tests.test_bts_contracts import assets
from tests.test_bts_workspace import make_legacy


def values(**changes):
    result = complete_values()
    result.update(contract_type='11', contract_derogation='none', contract_conclusion='2026-09-14',
        contract_start='2026-09-01', contract_end='2028-08-31', training_end='2028-06-30',
        rncp='41000', employer_idcc='1351', training_hours='1350', remote_hours='0',
        cfa_siret='84089988400026', cfa_uai='0831774C')
    result.update(changes)
    return result


class CalculationTests(unittest.TestCase):
    def test_exact_official_mos_values_and_conclusion_cutover(self):
        before = quote(values(contract_conclusion='2026-08-31'), {})
        after = quote(values(contract_conclusion='2026-09-01'), {})
        self.assertEqual(before['annual_cents'], 882800)
        self.assertEqual(after['annual_cents'], 876500)
        self.assertEqual(before['total_opco_cents'], 1765600)
        self.assertEqual(after['total_opco_cents'], 1753000)
        self.assertEqual(before['reference'], '2025-09')
        self.assertEqual(after['reference'], '2026-09')
        self.assertEqual([p['days'] for p in after['periods']], [365, 366])

    def test_branch_changes_rate_for_the_same_bts(self):
        self.assertEqual(quote(values(employer_idcc='1516'), {})['annual_cents'], 704500)
        self.assertEqual(quote(values(employer_idcc='1486'), {})['annual_cents'], 680000)
        self.assertEqual(quote(values(employer_idcc='2216'), {})['annual_cents'], 882800)
        self.assertEqual(quote(values(rncp='RNCP38362', employer_idcc='1486'), {})['annual_cents'], 600000)

    def test_old_rncp_alias_is_only_used_in_its_published_edition(self):
        self.assertEqual(quote(values(rncp='35393', contract_conclusion='2026-08-31'), {})['annual_cents'], 882800)
        with self.assertRaises(WorkspaceError):
            quote(values(rncp='35393'), {})

    def test_opco_published_370_day_example_and_inclusive_single_day(self):
        rows = periods(dt.date(2025, 8, 29), dt.date(2026, 9, 2), 800000)
        self.assertEqual(sum(p['days'] for p in rows), 370)
        self.assertEqual(sum(p['opco_cents'] for p in rows), 810959)
        self.assertEqual(periods(dt.date(2026, 9, 1), dt.date(2026, 9, 1), 800000)[0]['opco_cents'], 2192)

    def test_full_leap_year_is_one_annual_rate_and_partial_year_is_prorated(self):
        rows = periods(dt.date(2027, 9, 1), dt.date(2028, 8, 31), 800000)
        self.assertEqual(rows[0]['days'], 366)
        self.assertEqual(rows[0]['year_days'], 366)
        self.assertEqual(rows[0]['opco_cents'], 800000)
        rows = periods(dt.date(2027, 9, 1), dt.date(2028, 2, 29), 800000)
        self.assertEqual(rows[0]['days'], 182)
        self.assertEqual(rows[0]['opco_cents'], 397814)

    def test_distance_threshold_and_floor(self):
        self.assertFalse(quote(values(remote_hours='1079'), {})['remote_reduction'])
        distant = quote(values(remote_hours='1080'), {})
        self.assertEqual(distant['adjusted_annual_cents'], 701200)
        self.assertEqual(distant['total_opco_cents'], 1402400)
        custom = copy.deepcopy(reference()[-1])
        group = next(c for c in custom['certifications'] if '41000' in c['codes'])
        group['rates']['204'][0] = 450000
        self.assertEqual(quote(values(remote_hours='1350'), {}, [custom])['adjusted_annual_cents'], 400000)

    def test_missing_unknown_or_unsupported_inputs_never_become_zero_funding(self):
        for changes in ({'employer_idcc': ''}, {'employer_idcc': '9999'}, {'rncp': '999999'},
                        {'contract_conclusion': ''}, {'contract_conclusion': '2025-08-31'},
                        {'remote_hours': ''}, {'remote_hours': '1351'}, {'training_hours': 'NaN'},
                        {'contract_type': '36'}, {'employer_sector': 'public'},
                        {'contract_end': '2026-08-31'}, {'contract_end': '2030-08-31'}):
            with self.subTest(changes=changes):
                resolved = defaults(values(**changes), {})
                self.assertIn('error', resolved['_npec'])
                self.assertEqual(resolved['npec_1'], '')
                self.assertTrue(document_errors('formation', values(**changes), resolved))

    def test_ambiguous_idcc_requires_a_matching_branch(self):
        with self.assertRaises(WorkspaceError) as caught:
            quote(values(employer_idcc='3252'), {})
        self.assertEqual(len(caught.exception.cpne_choices), 3)
        result = quote(values(employer_idcc='3252'), {'npec_cpne': '433'})
        self.assertEqual(result['cpne'], '433')
        with self.assertRaises(WorkspaceError):
            quote(values(employer_idcc='1351'), {'npec_cpne': '433'})

    def test_server_derived_amounts_and_optional_company_balance(self):
        raw = validate_settings({'funding_mode': 'npec', 'npec_1': '999999.00', 'funding_years': '3', 'rac_1': '125,50'})
        result = defaults(values(), raw)
        self.assertEqual(result['npec_1'], '8765.00')
        self.assertEqual(result['funding_years'], '2')
        self.assertEqual(result['rac_1'], '125.50')
        self.assertEqual(result['rac_2'], '0.00')
        self.assertEqual(result['npec_3'], '')
        self.assertEqual(defaults(values(), result), result)

    def test_legacy_manual_settings_and_package_fingerprint_are_preserved(self):
        manual = {'funding_years': '2', 'npec_1': '8000', 'npec_2': '8000', 'rac_1': '0', 'rac_2': '100'}
        result = defaults(values(), manual)
        self.assertNotIn('_npec', result)
        self.assertNotIn('funding_mode', result)
        self.assertEqual(result['npec_1'], '8000')
        self.assertEqual(fingerprint(values(), result), fingerprint(values(), defaults(values(), result)))
        self.assertNotEqual(fingerprint(values(), result), fingerprint(values(), defaults(values(), dict(result, funding_mode='npec'))))

    def test_convention_uses_computed_annual_and_total_amounts(self):
        v = values()
        settings = defaults(v, {'teaching_mode': 'presentiel', 'employer_first_name': 'Alex', 'employer_last_name': 'EXEMPLE'})
        doc = Document(io.BytesIO(fill_convention('formation', v, settings, assets())))
        text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn('8 765,00', text)
        self.assertIn('17 530,00', text)
        self.assertNotIn('8 828,00', text)


class RouteTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.legacy = make_legacy(self.tmp.name); register_bts_workspace(self.legacy)
        self.store = ContractStore(self.legacy.AKTO_BTS_DB_FILE)
        self.rid = self.store.create_local({'apprentice_first_name': 'Test', 'apprentice_last_name': 'NPEC'})
        self.store.save_cerfa_complements(self.rid, values(), 0, source_version(self.store.record(self.rid)), 'Test')
        self.client = self.legacy.app.test_client()
        with self.client.session_transaction() as state:
            state.update(admin_logged_in=True, admin_role='admin')
        self.url = '/admin/BTS/dossiers/' + self.rid
        self.client.get(self.url + '?tab=contrat')
        with self.client.session_transaction() as state:
            self.csrf = {'bts_csrf_token': state['bts_csrf_token']}

    def test_quote_is_visible_without_writing_and_only_authorized_save_persists(self):
        page = self.client.get(self.url + '?tab=contrat')
        self.assertEqual(page.status_code, 200)
        self.assertIn('data-npec-financing', page.text)
        self.assertIn('8765.00', page.text)
        self.assertEqual(self.store.settings(self.rid)['revision'], 0)
        self.assertEqual(self.client.post(self.url + '/conventions/parametres', data={}).status_code, 400)
        payload = dict(self.csrf, revision='0', funding_mode='legacy', funding_years='3', npec_1='999999', rac_1='12.30')
        response = self.client.post(self.url + '/conventions/parametres', data=payload)
        self.assertEqual(response.status_code, 302)
        saved = self.store.settings(self.rid)
        self.assertEqual(saved['values']['funding_mode'], 'npec')
        self.assertEqual(saved['values']['npec_1'], '8765.00')
        self.assertEqual(saved['values']['rac_1'], '12.30')
        self.assertEqual(saved['values']['_npec']['reference'], '2026-09')
        self.client.post(self.url + '/conventions/parametres', data=dict(payload, rac_1='99'))
        self.assertEqual(self.store.settings(self.rid)['values']['rac_1'], '12.30')

    def test_old_document_snapshot_is_retained_after_an_explicit_recalculation(self):
        old = defaults(values(), {'npec_1': '8000', 'npec_2': '8000', 'funding_years': '2', 'rac_1': '0', 'rac_2': '0'})
        self.store.save_settings(self.rid, old, 0, 'Test')
        current = effective_values(self.store.record(self.rid), self.store.cerfa_complements(self.rid)['values'])
        package = {'id': 'npec-test-package', 'created_at': '2026-09-14T12:00:00Z', 'settings': old, 'fingerprint': fingerprint(current, old), 'documents': {}, 'signature': {}, 'opco': {}}
        self.store.save_package(self.rid, package)
        self.assertFalse(contract_context(self.legacy, self.store.record(self.rid))['contract_flow']['stale'])
        self.client.post(self.url + '/conventions/parametres', data=dict(old, **self.csrf, revision='1', funding_action='npec'))
        self.assertTrue(contract_context(self.legacy, self.store.record(self.rid))['contract_flow']['stale'])
        self.assertEqual(self.store.package(self.rid)['settings']['npec_1'], '8000')
        self.assertEqual(self.store.settings(self.rid)['values']['npec_1'], '8765.00')


if __name__ == '__main__':
    unittest.main()
