import copy
import json
import subprocess
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document

import app
import training_convocations as convocations


@pytest.fixture
def case(tmp_path, monkeypatch):
    trainee = {"id": "T-REFRESH", "public_token": "refresh-token", "first_name": "Camille",
               "last_name": "EXEMPLE", "email": "camille@example.test",
               "convocation_aps_sent_at": "2026-10-01T09:00:00Z", "convocation_aps_status": "sent",
               "convention_signature": {"status": "done", "signed_pdf_path": "signed-original.pdf"}}
    training = {"id": "S-REFRESH", "name": "A3P NOVEMBRE", "training_type": "A3P",
                "date_start": "2026-11-09", "date_end": "2027-01-19", "exam_date": "2027-01-20",
                "trainees": [trainee]}
    data = {"sessions": [training]}
    output = tmp_path / "generated_documents" / "convocations_aps"
    output.mkdir(parents=True)
    old = output / "issued.pdf"
    old.write_bytes(b"%PDF-1.4\nold dates: 09/11/2026 - 19/01/2027")
    trainee.update(convocation_aps_pdf_path=str(old),
                   convocation_aps_pdf_token="generated_documents/convocations_aps/issued.pdf")
    monkeypatch.setattr(app, "APS_CONVOCATION_DIR", str(output))
    monkeypatch.setattr(app, "PERSIST_DIR", str(tmp_path))
    monkeypatch.setattr(app, "load_data", lambda **kw: copy.deepcopy(data))
    save = Mock(side_effect=lambda mutate: mutate(data))
    monkeypatch.setattr(app, "_atomic_update_data", save)
    mail = Mock()
    monkeypatch.setattr(app, "brevo_send_email", mail)
    monkeypatch.setattr(app, "brevo_send_sms", mail)

    def generate(training, trainee, *ids):
        context = app._build_aps_convocation_context(training, trainee)
        revision = convocations.current_source_hash(app, training, trainee)
        path = output / ("convocation_" + revision + ".pdf")
        path.write_bytes(b"%PDF-1.4\n" + json.dumps(context).encode())
        trainee["convocation_aps_source_hash"] = revision
        return str(path.with_suffix(".docx")), str(path)

    generator = Mock(side_effect=generate)
    monkeypatch.setattr(app, "_generate_aps_convocation_files", generator)
    client = app.app.test_client()
    with client.session_transaction() as auth:
        auth["public_auth_refresh-token"] = True
    yield SimpleNamespace(trainee=trainee, training=training, data=data, output=output, old=old,
                          generate=generate, generator=generator, save=save, mail=mail, client=client)
    mail.assert_not_called()


URL = "/espace/refresh-token/convocation.pdf"


def test_legacy_a3p_convocation_refreshes_to_current_dates_once(case):
    case.training.update(date_start="2026-11-10", date_end="2027-01-14", exam_date="2027-01-15")
    original_signature = copy.deepcopy(case.trainee["convention_signature"])
    response = case.client.get(URL)
    assert response.status_code == 200
    assert b"10/11/2026" in response.data and b"14/01/2027" in response.data and b"15/01/2027" in response.data
    assert b"09/11/2026" not in response.data and b"19/01/2027" not in response.data
    assert case.trainee["convocation_aps_sent_at"] == "2026-10-01T09:00:00Z"
    assert case.trainee["convocation_aps_status"] == "sent"
    assert case.trainee["convention_signature"] == original_signature
    assert case.old.is_file()
    assert case.client.get(URL).data == response.data
    case.generator.assert_called_once()
    case.save.assert_called_once()


@pytest.mark.parametrize("target,key,value", [
    ("training", "date_start", "2026-11-10"), ("training", "date_end", "2027-01-14"),
    ("training", "exam_date", "2027-01-15"), ("training", "convocation_time", "09h00"),
    ("training", "exam_time", "09h30"), ("training", "name", "A3P SESSION MODIFIÉE"),
    ("trainee", "first_name", "Alex"), ("trainee", "address", "5 rue Exemple"),
])
def test_relevant_modification_refreshes_existing_pdf(case, target, key, value):
    first = case.client.get(URL).data
    getattr(case, target)[key] = value
    second = case.client.get(URL)
    assert second.status_code == 200 and second.data != first
    assert case.generator.call_count == 2
    assert case.trainee["convocation_aps_source_hash"] == convocations.current_source_hash(app, case.training, case.trainee)


def test_template_change_refreshes_pdf(case, tmp_path, monkeypatch):
    template = tmp_path / "template.docx"
    template.write_bytes(Path(app._aps_template_path(case.training)).read_bytes())
    monkeypatch.setattr(app, "_aps_template_path", lambda _: str(template))
    assert case.client.get(URL).status_code == 200
    first = case.trainee["convocation_aps_source_hash"]
    template.write_bytes(template.read_bytes() + b"updated-template")
    assert case.client.get(URL).status_code == 200
    assert case.trainee["convocation_aps_source_hash"] != first
    assert case.generator.call_count == 2


def test_old_download_links_keep_pointing_to_latest_pdf(case):
    legacy = case.trainee["convocation_aps_pdf_token"]
    case.client.get(URL)
    case.training["date_start"] = "2026-11-10"
    latest = case.client.get(URL)
    response = case.client.get("/espace/refresh-token/download/" + legacy + "?download=1", follow_redirects=True)
    assert response.status_code == 200 and response.data == latest.data
    assert response.headers["Content-Disposition"].startswith("attachment;")
    assert response.headers["Cache-Control"] == "private, no-store"
    assert case.generator.call_count == 2
    assert case.client.get("/espace/refresh-token/download/generated_documents/convocations_aps/other.pdf").status_code == 403


def test_admin_and_public_see_same_current_document(case):
    with case.client.session_transaction() as auth:
        auth["admin_logged_in"] = True
        auth["admin_role"] = "admin"
    case.training["date_start"] = "2026-11-10"
    response = case.client.get("/admin/sessions/S-REFRESH/stagiaires/T-REFRESH/convocation-aps.pdf")
    assert response.status_code == 200 and b"10/11/2026" in response.data
    assert "no-store" in response.headers["Cache-Control"]
    assert case.client.get(URL).data == response.data
    case.generator.assert_called_once()


def test_authenticated_access_required_before_refresh(case):
    with case.client.session_transaction() as auth:
        auth.clear()
    for url in [URL, "/espace/refresh-token/download/" + case.trainee["convocation_aps_pdf_token"]]:
        response = case.client.get(url, follow_redirects=True)
        assert response.request.path.endswith("/login")
    case.generator.assert_not_called()
    case.save.assert_not_called()


def test_unissued_convocation_is_not_created(case):
    for key in list(case.trainee):
        if key.startswith("convocation_aps_"):
            del case.trainee[key]
    assert case.client.get(URL).status_code == 404
    case.generator.assert_not_called()


def test_failed_refresh_never_serves_stale_pdf_or_changes_metadata(case):
    before = copy.deepcopy(case.trainee)
    case.generator.side_effect = RuntimeError("Conversion failed")
    response = case.client.get(URL)
    assert response.status_code == 302 and response.location.endswith("/espace/refresh-token")
    assert b"old dates" not in response.data
    assert case.trainee == before and case.old.is_file()
    case.save.assert_not_called()


def test_hash_for_another_generated_file_cannot_validate_an_old_pdf(case):
    # A generation followed by a failed email can persist the new hash while
    # the dossier still points to its previous file.
    case.trainee["convocation_aps_source_hash"] = convocations.current_source_hash(app, case.training, case.trainee)
    case.trainee["convocation_aps_source_pdf_path"] = str(case.output / "different-generated.pdf")
    response = case.client.get(URL)
    assert response.status_code == 200 and b"old dates" not in response.data
    case.generator.assert_called_once()


def test_cancelled_registration_keeps_archive_without_generating(case):
    case.trainee["registration_cancelled"] = True
    response = case.client.get(URL)
    assert response.status_code == 200 and response.data == case.old.read_bytes()
    case.generator.assert_not_called()
    case.save.assert_not_called()


def test_browser_conditionals_cannot_keep_old_convocation(case):
    first = case.client.get(URL)
    case.training["date_start"] = "2026-11-10"
    response = case.client.get(URL, headers={"If-None-Match": first.headers["ETag"],
                                            "If-Modified-Since": first.headers["Last-Modified"]})
    assert response.status_code == 200 and b"10/11/2026" in response.data
    assert response.headers["Cache-Control"] == "private, no-store"


def test_concurrent_date_change_prevents_publishing_obsolete_pdf(case):
    def generate_and_change(*args):
        result = case.generate(*args)
        case.training["date_start"] = "2026-11-10"
        return result
    case.generator.side_effect = generate_and_change
    response = case.client.get(URL)
    assert response.status_code == 302
    assert case.training["date_start"] == "2026-11-10"
    assert case.trainee["convocation_aps_pdf_path"] == str(case.old)
    assert "convocation_aps_source_hash" not in case.trainee


def test_refresh_preserves_concurrent_unrelated_changes(case):
    def generate_and_change(*args):
        result = case.generate(*args)
        case.trainee["notes"] = "Modification pendant la conversion"
        case.data["unrelated"] = {"value": 42}
        return result
    case.generator.side_effect = generate_and_change
    assert case.client.get(URL).status_code == 200
    assert case.trainee["notes"] == "Modification pendant la conversion"
    assert case.data["unrelated"] == {"value": 42}


def test_daily_issue_date_does_not_invalidate_current_document(tmp_path):
    template = tmp_path / "template.docx"
    template.write_bytes(b"template")
    first = convocations.source_hash({"date_jour": "27/10/2026", "date_convocation": "10/11/2026"}, template)
    second = convocations.source_hash({"date_jour": "28/10/2026", "date_convocation": "10/11/2026"}, template)
    assert first == second


def test_individual_ssiap_and_afc_dates_are_part_of_revision():
    trainee = {"first_name": "Camille", "last_name": "EXEMPLE", "email": "camille@example.test"}
    ssiap = {"name": "SSIAP NOVEMBRE", "training_type": "SSIAP", "date_start": "2026-11-09",
             "date_end": "2026-11-20", "exam_date": "2026-11-23"}
    ssiap["ssiap_pathways"] = app.ssiap_pathways.validate(ssiap, {
        **app.ssiap_pathways.DEFAULTS, "sst_dates": ["2026-11-05", "2026-11-06"],
        "ssiap_price": "980", "ssiap_sst_price": "1120"})
    assert convocations.current_source_hash(app, ssiap, {**trainee, "ssiap_pathway": "ssiap"}) != \
        convocations.current_source_hash(app, ssiap, {**trainee, "ssiap_pathway": "ssiap_sst"})
    afc = {"name": "AFC APS + SSIAP NOVEMBRE", "training_type": "APS", "date_start": "2026-11-16",
           "date_end": "2027-02-15", "exam_date": "2026-12-21", "ssiap_exam_date": "2027-02-16"}
    assert convocations.current_source_hash(app, afc, trainee) != \
        convocations.current_source_hash(app, afc, {**trainee, "training_date_start": "2026-11-23"})


def test_real_docx_render_and_failed_conversion_keep_previous_files(tmp_path, monkeypatch):
    template = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("{{date_debut_formation}} - {{date_fin_formation}} - {{date_examen}}")
    doc.save(template)
    monkeypatch.setattr(app, "_aps_template_path", lambda _: str(template))
    monkeypatch.setattr(app, "APS_CONVOCATION_DIR", str(tmp_path / "generated"))
    monkeypatch.setattr(app, "_find_libreoffice_binary", lambda: "libreoffice")
    trainee = {"id": "T-TEST", "first_name": "Camille", "last_name": "EXEMPLE", "email": "c@example.test"}
    trainee["convocation_aps_pdf_token"] = "previous-issued.pdf"
    training = {"training_type": "A3P", "date_start": "2026-11-10", "date_end": "2027-01-14", "exam_date": "2027-01-15"}

    def convert(command, **kwargs):
        staged = Path(command[-1])
        text = " ".join(p.text for p in Document(staged).paragraphs)
        assert text == "10/11/2026 - 14/01/2027 - 15/01/2027"
        assert any(value.startswith("-env:UserInstallation=") for value in command)
        staged.with_suffix(".pdf").write_bytes(b"%PDF-1.4\n" + text.encode())
    monkeypatch.setattr(app.subprocess, "run", convert)
    docx, pdf = app._generate_aps_convocation_files(training, trainee)
    assert trainee["convocation_aps_source_pdf_path"] == pdf
    assert "previous-issued.pdf" in trainee["convocation_aps_previous_pdf_tokens"]
    before = copy.deepcopy(trainee)
    original = Path(pdf).read_bytes()

    def fail(*args, **kwargs):
        raise subprocess.CalledProcessError(1, args[0], stderr="conversion failure")
    monkeypatch.setattr(app.subprocess, "run", fail)
    with pytest.raises(RuntimeError, match="conversion failure"):
        app._generate_aps_convocation_files(training, trainee)
    assert Path(pdf).read_bytes() == original and Path(docx).is_file()
    assert trainee == before
    assert not list(Path(app.APS_CONVOCATION_DIR).glob(".convocation-*"))
