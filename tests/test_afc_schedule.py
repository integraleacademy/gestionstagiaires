import copy
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document

import app
import afc_schedule


@pytest.fixture
def training():
    return {
        "id": "AFC-TEST", "name": "AFC APS + SSIAP NOVEMBRE 2026", "training_type": "APS",
        "partner_id": app.INTEGRALE_PARTNER_ID,
        "date_start": "2026-11-16", "date_end": "2027-02-15",
        "exam_date": "2026-12-21", "ssiap_exam_date": "2027-02-16",
        "trainees": [],
    }


@pytest.fixture
def trainee():
    return {"id": "TRN-AFC-TEST", "first_name": "Camille", "last_name": "MARTIN",
            "email": "camille@example.test", "documents": [],
            "training_date_start": "2026-11-23", "training_date_end": "2027-02-16"}


@pytest.fixture
def client():
    client = app.app.test_client()
    with client.session_transaction() as session:
        session["admin_logged_in"] = True
        session["admin_role"] = "admin"
    return client


def test_afc_does_not_compute_or_move_aps_exam(training):
    before = copy.deepcopy(training)
    app._sync_aps_period_dates(training)
    assert training == before
    assert afc_schedule.enabled(training)
    assert not afc_schedule.enabled({"name": "APS novembre", "training_type": "APS"})


def test_create_and_partial_update_use_one_period_and_two_exams(client, training):
    data = {"sessions": []}
    payload = {k: v for k, v in training.items() if k not in {"id", "trainees"}}
    with patch.object(app, "load_data", return_value=data), patch.object(app, "save_data"):
        created = client.post("/api/sessions/create", json=payload)
        assert created.status_code == 200
        saved = data["sessions"][0]
        for key in afc_schedule.SESSION_DATES:
            assert saved[key] == payload[key]
        saved["ssiap_exam"] = {"config": {"date": "2027-02-16"}, "revision": 1}
        result = client.post(f"/api/sessions/{saved['id']}/update", json={"ssiap_exam_date": "2027-02-17"})
        assert result.status_code == 200
        assert saved["date_end"] == "2027-02-15"
        assert saved["exam_date"] == "2026-12-21"
        assert saved["date_start"] == "2026-11-16"
        assert not any(key.startswith(("afc_aps_", "afc_ssiap_")) for key in saved)
        assert saved["ssiap_exam"]["config"]["date"] == "2027-02-17"
        assert saved["ssiap_exam"]["revision"] == 2


@pytest.mark.parametrize("payload", [
    {"date_end": "2026-11-01"}, {"exam_date": "2026-12-32"},
    {"ssiap_exam_date": "2027-02-30"}, {"date_start": "2027-02-30"},
    {"exam_date": ["2026-12-21"]},
])
def test_session_validation_is_atomic(client, training, payload):
    before = copy.deepcopy(training)
    with patch.object(app, "load_data", return_value={"sessions": [training]}), patch.object(app, "save_data") as save:
        response = client.post(f"/api/sessions/{training['id']}/update", json=payload)
        assert response.status_code == 400
        assert training == before
        save.assert_not_called()


def test_student_dates_save_clear_and_do_not_change_another_student(client, training, trainee):
    other = {"id": "TRN-OTHER", "first_name": "Alex", "documents": []}
    training["trainees"] = [trainee, other]
    original_session_dates = (training["date_start"], training["date_end"])
    url = f"/api/sessions/{training['id']}/stagiaires/{trainee['id']}/update"
    with patch.object(app, "load_data", return_value={"sessions": [training]}), patch.object(app, "save_data"), \
         patch.object(app, "brevo_send_email") as mail:
        response = client.post(url, json={"training_date_start": "2026-11-30", "training_date_end": "2027-02-16"})
        assert response.status_code == 200, response.json
        assert trainee["training_date_start"] == "2026-11-30"
        assert app._trainee_training_session(training, other)["date_start"] == "2026-11-16"
        response = client.post(url, json={"training_date_start": "", "training_date_end": ""})
        assert response.status_code == 200
        assert app._trainee_training_session(training, trainee)["date_start"] == "2026-11-16"
        assert app._trainee_training_session(training, trainee)["date_end"] == "2027-02-15"
        mail.assert_not_called()
    assert (training["date_start"], training["date_end"]) == original_session_dates


def test_reversed_individual_period_is_rejected_without_save(client, training, trainee):
    training["trainees"] = [trainee]
    before = copy.deepcopy(trainee)
    with patch.object(app, "load_data", return_value={"sessions": [training]}), patch.object(app, "save_data") as save:
        response = client.post(f"/api/sessions/{training['id']}/stagiaires/{trainee['id']}/update",
                               json={"training_date_start": "2027-03-01"})
    assert response.status_code == 400
    assert trainee == before
    save.assert_not_called()


def test_new_trainee_dates_are_saved_before_any_document_automation(client, training):
    with patch.object(app, "load_data", return_value={"sessions": [training]}), \
         patch.object(app, "save_data"), patch.object(app, "brevo_send_email") as mail:
        response = client.post(f"/api/sessions/{training['id']}/trainees/create", json={
            "first_name": "Camille", "last_name": "MARTIN", "email": "camille@example.test", "send_access": False,
            "training_date_start": "2026-11-23", "training_date_end": "2027-02-16",
        })
    assert response.status_code == 200, response.json
    saved = training["trainees"][0]
    assert saved["training_date_start"] == "2026-11-23"
    assert saved["training_date_end"] == ""
    assert app._build_aps_convocation_context(training, saved)["date_convocation"] == "23/11/2026"
    mail.assert_not_called()


def test_documents_use_individual_dates_and_both_exams_without_mutation(training, trainee, tmp_path):
    before = copy.deepcopy(training)
    context = app._build_aps_convocation_context(training, trainee)
    convention = app._aps_convention_replacements(training, trainee)
    for values in (context, convention):
        assert values["date_debut_formation"] == "23/11/2026"
        assert values["date_fin_formation"] == "15/02/2027"
        assert values["date_convocation"] == "23/11/2026"
        assert values["date_examen_aps"] == "21/12/2026"
        assert values["date_examen_ssiap"] == "16/02/2027"
        assert values["periode_formation"] == "du 23/11/2026 au 15/02/2027"
        assert "periode_aps" not in values
        assert "periode_ssiap" not in values
    assert training == before
    assert app._automation_document_config(training)["slug"] == "afc"
    assert app._can_send_convocation_without_signed_convention(training)
    # Render the actual business placeholders in both stored models.
    for name, values in (("conventionafc.docx", convention), ("convocationafc.docx", context)):
        import shutil
        out = tmp_path / name
        shutil.copyfile(Path("templates_word") / name, out)
        app._replace_docx_xml_placeholders(str(out), values)
        paragraphs = list(app._iter_docx_paragraphs(Document(out)))
        text = "\n".join(p.text for p in paragraphs)
        for date in ("23/11/2026", "16/02/2027", "21/12/2026"):
            assert date in text
        if name.startswith("convention"):
            remaining, anchors = app._assert_aps_convention_has_no_business_placeholders(str(out))
            assert anchors
        else:
            app._assert_docx_has_no_unresolved_variables(str(out))


def test_email_and_document_revision_follow_individual_calendar(training, trainee):
    subject, body = app._build_aps_convocation_email("Camille", training["date_start"], training["date_end"], training, trainee)
    assert "APS + SSIAP 1" in subject
    for date in ("23/11/2026", "16/02/2027", "21/12/2026"):
        assert date in body
    revision = afc_schedule.revision(training, trainee)
    trainee["training_date_start"] = "2026-11-30"
    assert revision != afc_schedule.revision(training, trainee)


def test_missing_exam_date_is_reported_before_document_generation(training, trainee):
    del training["ssiap_exam_date"]
    with pytest.raises(ValueError, match="Examen SSIAP 1"):
        app._build_aps_convocation_context(training, trainee)
    with pytest.raises(ValueError, match="Examen SSIAP 1"):
        app._aps_convention_replacements(training, trainee)


def test_retired_component_dates_do_not_constrain_or_appear_in_documents(client, training, trainee):
    training.update(afc_aps_start="2027-04-01", afc_aps_end="2026-10-01",
                    afc_ssiap_start="2027-05-01", afc_ssiap_end="2027-04-01")
    with patch.object(app, "load_data", return_value={"sessions": [training]}), patch.object(app, "save_data"):
        response = client.post(f"/api/sessions/{training['id']}/update", json={"exam_date": "2026-12-01"})
    assert response.status_code == 200
    for values in (app._build_aps_convocation_context(training, trainee), app._aps_convention_replacements(training, trainee)):
        assert values["date_examen_aps"] == "01/12/2026"
        assert values["periode_formation"] == "du 23/11/2026 au 15/02/2027"
        assert "01/04/2027" not in str(values)


def test_session_page_exposes_afc_fields_and_does_not_change_dates(client, training):
    before = copy.deepcopy(training)
    with patch.object(app, "load_data", return_value={"sessions": [training]}), \
         patch.object(app, "_load_wedof_webhooks", return_value=[]):
        response = client.get("/admin/sessions")
    assert response.status_code == 200
    text = response.get_data(as_text=True)
    for label in ("Examen APS", "Examen SSIAP 1", "Début formation", "Fin formation"):
        assert label in text
    for retired in ("AfcApsStart", "AfcApsEnd", "AfcSsiapStart", "AfcSsiapEnd", "afc_aps_", "afc_ssiap_"):
        assert retired not in text
    assert 'data-afc-enabled="1"' in text
    assert training["exam_date"] == before["exam_date"]
    assert training["date_end"] == before["date_end"]
