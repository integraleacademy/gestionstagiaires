"""Fill the supplied Word conventions and prepare individual PDFs for signing."""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from pathlib import Path
from lxml import etree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ContentStream, NameObject, TextStringObject, ByteStringObject
from pypdf._cmap import build_char_map
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

from bts_cerfa import readiness
from bts_cerfa_pdf import generate_pdf
from bts_workspace_store import WorkspaceError, now, money_cents, date_fr
from bts_npec import resolve as resolve_financing

ROOT = Path(__file__).resolve().parent
TEMPLATES = ROOT / "templates_word" / "bts"
VERSION = "20260914-1"
DOC_LABELS = {"cerfa": "Contrat d’apprentissage", "formation": "Convention de formation",
              "mobilite": "Convention de mobilité"}
MODES = (("presentiel", "En présentiel"), ("distanciel", "À distance"), ("hybride", "Hybride"))
MOS_GOALS = (
    "Le titulaire du BTS « Management opérationnel de la sécurité » exerce ses missions dans le cadre de la sécurité et de la sûreté. "
    "Il conduit ses activités, soit dans une entreprise de sécurité/sûreté, soit dans le service interne de sécurité d'une organisation "
    "(entreprise, administration publique et privée, association). En outre, il participe à l'organisation du service dans lequel il évolue. "
    "Il est donc à même de manager des équipes, de participer à la gestion administrative et juridique du personnel. "
    "Il est en relation directe avec le client et en contact permanent avec sa hiérarchie. Enfin, il assure la liaison avec les différents "
    "acteurs institutionnels et participe aux instances de sécurité.")
MOS_CONTENT = "\n".join(("Culture générale et expression", "Langue vivante (anglais)", "Culture économique, juridique et managériale",
    "Préparation et mise en œuvre d’une prestation de sécurité", "Management des ressources humaines",
    "Gestion de la relation client", "Participation à la sécurité globale"))

SETTINGS_FIELDS = (
    ("teaching_mode", "Mode de formation", "mode"),
    ("convention_date", "Date d’établissement des conventions", "date"),
    ("employer_first_name", "Prénom du signataire de l’entreprise", "text"),
    ("employer_last_name", "Nom du signataire de l’entreprise", "text"),
    ("employer_signer_email", "E-mail du signataire de l’entreprise", "email"),
    ("guardian_first_name", "Prénom du représentant légal", "text"),
    ("guardian_last_name", "Nom du représentant légal", "text"),
    ("guardian_phone", "Téléphone du représentant légal", "tel"),
    ("funding_mode", "Calcul du financement", "funding"),
    ("npec_cpne", "Branche professionnelle", "text"),
    ("funding_years", "Nombre d’années de financement", "years"),
    *((f"{prefix}_{year}", f"Année {year} · {label} (€)", "money") for year in range(1, 4)
      for prefix, label in (("npec", "Prise en charge OPCO"), ("rac", "Reste à charge entreprise"))),
    ("meal_price", "Restauration · montant par repas (€)", "money"),
    ("meals_1", "Restauration · nombre de repas année 1", "integer"),
    ("meals_2", "Restauration · nombre de repas année 2", "integer"),
    ("equipment_cost", "Premier équipement pédagogique (€)", "money"),
    ("mobility_cost", "Mobilité · montant (€)", "money"),
    ("mobility_start", "Mobilité · date de début", "date"),
    ("mobility_end", "Mobilité · date de fin", "date"),
    ("training_rhythm", "Rythme d’alternance", "text"),
    ("training_goals", "Objectifs de la formation", "textarea"),
    ("training_content", "Programme de la formation", "textarea"),
    ("financer", "OPCO destinataire", "financer"),
)


def is_mos(values):
    title = values.get("training_title", "").casefold()
    return "mos" in title or "sécurité" in title or "securite" in title


def defaults(values, saved):
    result = {"convention_date": dt.date.today().isoformat(), "funding_years": "2",
              "meal_price": "3", "meals_1": "100", "meals_2": "100", "equipment_cost": "500",
              "mobility_cost": "1500", "mobility_start": "2027-11-08", "mobility_end": "2027-11-12",
              "employer_signer_email": values.get("employer_email", "")}
    if is_mos(values):
        result.update(training_goals=MOS_GOALS, training_content=MOS_CONTENT,
                      training_rhythm="2 semaines à l’école – 2 semaines en entreprise")
    result.update(saved)
    return resolve_financing(values, result)


def validate_settings(raw):
    raw = dict(raw)
    if raw.get("funding_action") == "npec":
        raw["funding_mode"] = "npec"
    if raw.get("funding_mode") == "npec":
        # The browser cannot supply or override any computed OPCO amount.
        for key in ("funding_years", "npec_1", "npec_2", "npec_3"):
            raw[key] = ""
    result = {}
    for key, label, kind in SETTINGS_FIELDS:
        value = str(raw.get(key) or "").strip()
        if len(value) > (6000 if kind == "textarea" else 250):
            raise WorkspaceError(f"Le champ « {label} » est trop long.")
        if value and kind == "date":
            try:
                dt.date.fromisoformat(value)
            except ValueError:
                raise WorkspaceError(f"La date « {label} » est invalide.") from None
        if value and kind in {"money", "integer"}:
            amount = money_cents(value, strict=True)
            if amount < 0 or (kind == "integer" and amount % 100):
                raise WorkspaceError(f"Le montant ou nombre « {label} » est invalide.")
            value = str(amount // 100) if kind == "integer" else f"{amount / 100:.2f}"
        if value and kind == "email" and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value):
            raise WorkspaceError(f"L’e-mail « {label} » est invalide.")
        if kind == "mode" and value not in {"", *(k for k, _ in MODES)}:
            raise WorkspaceError("Choisissez le mode de formation.")
        if kind == "years" and value not in {"", "1", "2", "3"}:
            raise WorkspaceError("Choisissez entre une et trois années de financement.")
        if kind == "funding" and value not in {"", "npec", "legacy"}:
            raise WorkspaceError("Mode de calcul du financement invalide.")
        if key == "npec_cpne" and value and not re.fullmatch(r"\d{1,5}", value):
            raise WorkspaceError("Branche professionnelle invalide.")
        if kind == "financer":
            from wedof_bts import FINANCERS
            if value and value not in FINANCERS:
                raise WorkspaceError("Choisissez un OPCO pris en charge.")
        result[key] = value
    if result.get("mobility_start") and result.get("mobility_end"):
        days = (dt.date.fromisoformat(result["mobility_end"]) - dt.date.fromisoformat(result["mobility_start"])).days + 1
        if not 1 <= days <= 28:
            raise WorkspaceError("Le modèle de mobilité est prévu pour une période de 1 à 28 jours.")
    return result


def needs_guardian(values, today=None):
    if not values.get("apprentice_birth_date"):
        raise WorkspaceError("Renseignez la date de naissance pour déterminer les signataires.")
    born = dt.date.fromisoformat(values["apprentice_birth_date"])
    today = today or dt.date.today()
    conclusion = dt.date.fromisoformat(values.get("contract_conclusion") or today.isoformat())
    reference = min(today, conclusion)
    age = reference.year - born.year - ((reference.month, reference.day) < (born.month, born.day))
    if age < 0:
        raise WorkspaceError("Vérifiez la date de naissance de l’apprenti.")
    return age < 18 and values.get("apprentice_emancipated") != "yes"


def document_errors(kind, values, settings):
    labels = {"apprentice_first_name": "Prénom de l’apprenti", "apprentice_last_name": "Nom de l’apprenti",
        "employer_name": "Raison sociale", "employer_address": "Adresse de l’entreprise", "employer_postcode": "Code postal de l’entreprise",
        "employer_city": "Ville de l’entreprise", "employer_siret": "SIRET", "training_title": "Intitulé du BTS",
        "rncp": "Code RNCP", "diploma_code": "Code diplôme", "training_start": "Début de formation", "training_end": "Fin de formation",
        "training_hours": "Durée de formation", "contract_start": "Début du contrat", "contract_end": "Fin du contrat"}
    required = list(labels) if kind == "formation" else list(labels)[:7]
    missing = [labels[k] for k in required if not values.get(k)]
    if kind == "formation":
        if settings.get("_npec", {}).get("error"):
            missing.append(settings["_npec"]["error"])
        if values.get("cfa_siret") != "84089988400026" or values.get("cfa_uai") != "0831774C":
            missing.append("Le CFA doit correspondre au modèle Intégrale Academy (SIRET 84089988400026, UAI 0831774C)")
    setting_keys = ["teaching_mode", "convention_date"]
    if kind == "formation":
        setting_keys += ["funding_years", "employer_first_name", "employer_last_name", "training_rhythm", "training_goals", "training_content", "meal_price", "meals_1", "equipment_cost"]
        for year in range(1, int(settings.get("funding_years") or 2) + 1):
            setting_keys += [f"npec_{year}", f"rac_{year}"]
        if int(settings.get("funding_years") or 2) > 1:
            setting_keys.append("meals_2")
    if settings.get("teaching_mode") == "presentiel":
        setting_keys += ["mobility_start", "mobility_end", "mobility_cost"]
    if kind == "mobilite" and settings.get("teaching_mode") != "presentiel":
        missing.append("La convention de mobilité concerne uniquement le présentiel")
    names = {k: label for k, label, _ in SETTINGS_FIELDS}
    missing += [names[k] for k in setting_keys if settings.get(k, "") == ""]
    if kind == "mobilite" and values.get("training_start") and values.get("training_end") and settings.get("mobility_start") and settings.get("mobility_end"):
        if settings["mobility_start"] < values["training_start"] or settings["mobility_end"] > values["training_end"]:
            missing.append("Les dates de mobilité doivent se situer pendant la formation")
    return missing


def signature_errors(values, settings):
    errors = []
    try:
        minor = needs_guardian(values)
    except (WorkspaceError, ValueError) as exc:
        errors.append(str(exc)); minor = False
    for prefix in ("employer",) + (("guardian",) if minor else ()):
        for part in ("first_name", "last_name"):
            if not settings.get(prefix + "_" + part):
                errors.append(("Entreprise" if prefix == "employer" else "Représentant légal") + " : " + ("prénom" if part == "first_name" else "nom"))
    for email, label in ((values.get("apprentice_email"), "Apprenti"), (settings.get("employer_signer_email"), "Entreprise"),
                         *(([(values.get("guardian_email"), "Représentant légal")]) if minor else [])):
        if not email or not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            errors.append(label + " : adresse e-mail")
    return errors


def fingerprint(values, settings):
    return hashlib.sha256(json.dumps({"version": VERSION, "values": values, "settings": settings},
                         ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def tuition(settings):
    return sum((money_cents(settings.get(f"npec_{n}"), strict=True) + money_cents(settings.get(f"rac_{n}"), strict=True))
               for n in range(1, int(settings["funding_years"]) + 1))


def money(value):
    return f"{value / 100:,.2f}".replace(",", " ").replace(".", ",")


def xml_replace(paragraph, replacements):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    nodes = paragraph.findall(".//w:t", ns)
    text = "".join(n.text or "" for n in nodes)
    matches = []
    for key, value in replacements.items():
        matches.extend((m.start(), m.end(), str(value)) for m in re.finditer(re.escape(key), text))
    for start, end, value in sorted(matches, reverse=True):
        pos = 0
        inserted = False
        for node in nodes:
            raw = node.text or ""
            stop = pos + len(raw)
            if stop > start and pos < end:
                before, after = raw[:max(0, start - pos)], raw[max(0, end - pos):]
                node.text = before + (value if not inserted else "") + after
                node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                inserted = True
            pos = stop


def paragraphs(container):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from paragraphs(cell)


def after_paragraph(paragraph, text):
    from docx.text.paragraph import Paragraph
    node = OxmlElement("w:p")
    paragraph._p.addnext(node)
    p = Paragraph(node, paragraph._parent)
    p.add_run(text)
    return p


def fill_convention(kind, values, settings, assets):
    errors = document_errors(kind, values, settings)
    if errors:
        raise WorkspaceError("À compléter : " + "; ".join(errors))
    replacements = {"NOM_ENTREPRISE": values.get("employer_name", ""), "SIRET_ENTREPRISE": values.get("employer_siret", ""),
        "ADR1_ENTREPRISE": values.get("employer_address", ""), "ADR2_ENTREPRISE": values.get("employer_address_complement", ""),
        "ADR3_ENTREPRISE": "", "ADR4_ENTREPRISE": "", "CP_ENTREPRISE": values.get("employer_postcode", ""),
        "VILLE_ENTREPRISE": values.get("employer_city", ""),
        "CIVILITE_APPRENANT": {"M": "Monsieur", "F": "Madame"}.get(values.get("apprentice_sex"), ""),
        "NOM_PRENOM_APPRENANT": " ".join(values.get(k, "") for k in ("apprentice_last_name", "apprentice_first_name")),
        "NOM_RESPONSABLE_ENTREPRISE": settings.get("employer_last_name", ""), "PRENOM_RESPONSABLE_ENTREPRISE": settings.get("employer_first_name", ""),
        "DATE_DEB_FORMATION_CERFA": date_fr(values.get("training_start")), "DATE_FIN_FORMATION_CERFA": date_fr(values.get("training_end")),
        "DUREE_HEURES_FORMATION": values.get("training_hours", ""), "DATE_DEB_CONTRAT_APPRENANT": date_fr(values.get("contract_start")),
        "DATE_FIN_CONTRAT_APPRENANT": date_fr(values.get("contract_end")), "DATE_ETABLISSEMENT_CONVENTION": date_fr(settings["convention_date"])}
    for n in range(1, 4):
        replacements[f"MONTANT_NPEC_ANNEE_{n}"] = money(money_cents(settings[f"npec_{n}"])) if settings.get(f"npec_{n}") else "Sans objet"
        replacements[f"MONTANT_RESTE_A_CHARGE_ANNEE_{n}"] = money(money_cents(settings[f"rac_{n}"])) if settings.get(f"rac_{n}") else "Sans objet"
    if kind == "formation":
        replacements["MONTANT_NPEC_TOTAL"] = money(tuition(settings))
    replacements = {"«" + k + "»": v for k, v in replacements.items()}
    replacements.update({"[sc_sign1.signature]": "BTSMARK_APPRENTICE" if kind == "mobilite" else "BTSMARK_EMPLOYER",
                         "[/sc_sign1.signature]": "", "[sc_sign2.signature]": "BTSMARK_EMPLOYER", "[/sc_sign2.signature]": "",
                         "[sc_user.signature]": "BTSMARK_CFA", "[/sc_user.signature]": ""})
    buffer = io.BytesIO()
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(TEMPLATES / f"convention_{kind}.dotx") as source, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml")
            if name.startswith("word/") and name.endswith(".xml"):
                root = ET.fromstring(data)
                for parent in root.iter():
                    for child in list(parent):
                        if child.tag in {f"{{{ns['w']}}}instrText", f"{{{ns['w']}}}fldChar"}:
                            parent.remove(child)
                for paragraph in root.findall(".//w:p", ns):
                    xml_replace(paragraph, replacements)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(name, data)
    buffer.seek(0)
    doc = Document(buffer)
    for p in list(paragraphs(doc)):
        if "BTSMARK_CFA" in p.text:
            p.clear()
            p.add_run().add_picture(io.BytesIO(assets["stamp"]), width=Inches(1.5))
            p.add_run().add_picture(io.BytesIO(assets["signature"]), width=Inches(.7))
        elif "BTSMARK_" in p.text:
            p.paragraph_format.space_after = Pt(46)
            for run in p.runs:
                run.font.size = Pt(5)
    if kind == "formation":
        # The source leaves several empty paragraphs after each
        # signature table. With a real stamp they create otherwise blank pages.
        from docx.oxml.ns import qn
        for table in doc.tables:
            if "BTSMARK_EMPLOYER" not in table._element.xml:
                continue
            for row in table.rows:
                for cell in row.cells:
                    for p in list(cell.paragraphs):
                        if len(cell.paragraphs) > 1 and not p.text.strip() and not p._p.findall(".//" + qn("w:drawing")) and not p._p.findall(".//" + qn("w:pict")):
                            p._p.getparent().remove(p._p)
            following = table._element.getnext()
            while following is not None and following.tag == qn("w:p") and not "".join(following.itertext()).strip() and not following.findall(".//" + qn("w:drawing")):
                next_node = following.getnext()
                following.getparent().remove(following)
                following = next_node
        all_paragraphs = list(paragraphs(doc))
        for p in all_paragraphs:
            text = p.text.strip()
            if text == "Annexe à la convention de formation":
                for br in p._p.findall(".//" + qn("w:br")):
                    if br.get(qn("w:type")) == "page":
                        br.getparent().remove(br)
                p.paragraph_format.page_break_before = True
            replacements2 = {"32034401": values["diploma_code"], "41000": re.sub(r"^RNCP\s*", "", values["rncp"], flags=re.I),
                "MANAGEMENT OPÉRATIONNEL DE LA SÉCURITÉ (MOS)": values["training_title"],
                "BTS MANAGEMENT OPÉRATIONNEL DE LA SÉCURITÉ": values["training_title"],
                "2 semaines au centre de formation – 2 semaines en entreprise": settings["training_rhythm"],
                "En présentiel": dict(MODES)[settings["teaching_mode"]],
                "en 2028": "en " + values["training_end"][:4],
                "0831774": "0831774C"}
            for key, value in replacements2.items():
                if key in p.text:
                    # Work directly on the OOXML run nodes to retain template formatting.
                    xml_replace(p._p, {key: value})
            if text.startswith("Le titulaire du BTS"):
                p.text = settings["training_goals"]
            if text in MOS_CONTENT.splitlines():
                if text == MOS_CONTENT.splitlines()[0]:
                    from copy import deepcopy
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    lines = [line.strip() for line in settings["training_content"].splitlines() if line.strip()]
                    p.text = lines[0]
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    previous = p
                    for line in lines[1:]:
                        inserted = after_paragraph(previous, line)
                        if p._p.pPr is not None:
                            inserted._p.insert(0, deepcopy(p._p.pPr))
                        previous = inserted
                else:
                    p._p.getparent().remove(p._p)
            if not is_mos(values) and text.startswith("Constituée de professionnels de la sécurité"):
                p.text = "Constituée de professionnels du métier pour les matières professionnelles et d’enseignants pour les matières générales."
            if "était sous statut de stagiaire" in text and values["contract_start"] <= values["training_start"]:
                p.text = "Sans objet : le contrat débute au plus tard à l’entrée en formation."
            if text == "Montant : 1500€":
                p.text = "Montant : " + (money(money_cents(settings["mobility_cost"])) if settings["teaching_mode"] == "presentiel" else "0,00") + " €"
            if text == "Montant : 500€":
                p.text = "Montant : " + money(money_cents(settings["equipment_cost"])) + " €"
        finance = doc.tables[3]
        for year in range(1, 4):
            row = finance.rows[year + 1]
            if year <= int(settings["funding_years"]):
                row.cells[1].text = money(money_cents(settings[f"npec_{year}"]) + money_cents(settings[f"rac_{year}"])) + " €"
            else:
                for cell in row.cells[1:]:
                    cell.text = "Sans objet"
        meals = doc.tables[5]
        meals.rows[1].cells[1].text = money(money_cents(settings["meal_price"])) + " € par repas"
        for year in (1, 2):
            count = int(settings.get(f"meals_{year}") or 0) if year <= int(settings["funding_years"]) else 0
            meals.rows[year + 1].cells[1].text = "Nbre : " + str(count)
            meals.rows[year + 1].cells[2].text = money(count * money_cents(settings["meal_price"])) + " €"
        # Only presentiel learners are assigned the London mobility in this workflow.
        if settings["teaching_mode"] != "presentiel":
            found = False
            for p in doc.paragraphs:
                if p.text.startswith("5.3"):
                    found = True
                elif found and p.text.strip() == "OUI":
                    p.text = "NON"; break
    else:
        for p in list(doc.paragraphs):
            if p.text.startswith("La présente convention s'applique du"):
                days = (dt.date.fromisoformat(settings["mobility_end"]) - dt.date.fromisoformat(settings["mobility_start"])).days + 1
                p.text = f"La présente convention s’applique du {date_fr(settings['mobility_start'])} au {date_fr(settings['mobility_end'])}, soit une durée totale de {days} jours."
            if p.text.startswith("Fait en 4 exemplaires"):
                p.text = "Fait à Puget sur Argens, le " + date_fr(settings["convention_date"]) + "."
            if "BTSMARK_APPRENTICE" in p.text and needs_guardian(values):
                p2 = after_paragraph(p, "Pour le représentant légal : " + " ".join(settings.get(k, "") for k in ("guardian_first_name", "guardian_last_name")))
                p2.paragraph_format.space_before = Pt(10)
                marker = after_paragraph(p2, "BTSMARK_GUARDIAN")
                marker.runs[0].font.size = Pt(5)
                marker.paragraph_format.space_after = Pt(46)
    for p in paragraphs(doc):
        if re.search(r"«[A-Z_0-9]+»|\[/?sc_", p.text):
            raise WorkspaceError("Une variable du modèle n’a pas été remplie.")
    output = io.BytesIO(); doc.save(output)
    return output.getvalue()


def convert_pdf(docx_path):
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    if not binary:
        raise WorkspaceError("La conversion PDF n’est pas disponible sur le serveur.")
    with tempfile.TemporaryDirectory(prefix="bts-office-") as profile:
        try:
            result = subprocess.run([binary, "-env:UserInstallation=" + Path(profile).as_uri(), "--headless", "--convert-to", "pdf",
                "--outdir", str(docx_path.parent), str(docx_path)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
        except (subprocess.TimeoutExpired, OSError):
            raise WorkspaceError("La conversion de la convention a échoué. Réessayez.") from None
    path = docx_path.with_suffix(".pdf")
    if result.returncode or not path.is_file():
        raise WorkspaceError("La convention n’a pas pu être convertie en PDF.")
    return path


def extract_signature_fields(pdf_path):
    """Strip dedicated marker text operators, preserving the PDF page geometry."""
    reader = PdfReader(pdf_path)
    writer, fields = PdfWriter(), []
    for number, page in enumerate(reader.pages, 1):
        height, width = float(page.mediabox.height), float(page.mediabox.width)
        def visitor(text, cm, tm, font, size):
            for match in re.finditer(r"BTSMARK_(EMPLOYER|APPRENTICE|GUARDIAN)", text or ""):
                # LibreOffice outputs an unrotated text matrix for these standalone markers.
                x, baseline = float(cm[4]) + float(tm[4]), float(cm[5]) + float(tm[5])
                y = height - baseline - float(size)
                if x < 0 or y < 0 or x + 150 > width or y + 45 > height:
                    raise WorkspaceError("Une zone de signature sort de la page. Vérifiez le modèle.")
                fields.append({"role": match.group(1).lower(), "type": "signature", "page": number,
                               "x": round(x), "y": round(y), "width": 150, "height": 45})
        page.extract_text(visitor_text=visitor)
        stream = ContentStream(page.get_contents(), reader)
        # LibreOffice uses subset fonts: raw strings contain glyph codes, not
        # Unicode. Decode with the page's actual font map before removing the
        # dedicated marker operator; other text and images remain untouched.
        kept, font, maps = [], None, {}
        for args, op in stream.operations:
            if op == b"Tf":
                font = str(args[0])
            decoded = ""
            if op in {b"Tj", b"TJ"} and font:
                if font not in maps:
                    maps[font] = build_char_map(font, 200, page)
                encoding, mapping = maps[font][2:4]
                for item in args[0] if op == b"TJ" else args:
                    if isinstance(item, (TextStringObject, ByteStringObject)):
                        raw = item.original_bytes if isinstance(item, TextStringObject) else bytes(item)
                        chars = raw.decode(encoding) if isinstance(encoding, str) else "".join(encoding.get(b, chr(b)) for b in raw)
                        decoded += "".join(mapping.get(c, c) for c in chars)
            if "BTSMARK_" not in decoded:
                kept.append((args, op))
        stream.operations = kept
        page[NameObject("/Contents")] = stream
        writer.add_page(page)
    if not fields:
        raise WorkspaceError("Aucune zone de signature n’a été trouvée dans la convention.")
    with pdf_path.open("wb") as handle:
        writer.write(handle)
    if any("BTSMARK_" in p.extract_text() for p in PdfReader(pdf_path).pages):
        raise WorkspaceError("Une zone de signature n’a pas pu être préparée correctement.")
    return fields


def prepare_cerfa(values, assets):
    reader = PdfReader(io.BytesIO(generate_pdf(values)))
    # Flatten field appearances before signing. Removing widgets without painting
    # their appearance streams would silently erase all entered CERFA values.
    writer = PdfWriter(clone_from=reader)
    for page in writer.pages:
        for ref in list(page.get("/Annots", [])):
            widget = ref.get_object()
            ap = widget.get("/AP", {}).get("/N")
            if not ap:
                continue
            ap = ap.get_object()
            if not hasattr(ap, "get_data"):
                ap = ap.get(widget.get("/AS", "/Off"))
                if not ap:
                    continue
                ap = ap.get_object()
            from pypdf.generic import DictionaryObject, RectangleObject, DecodedStreamObject
            rect = widget.get("/Rect")
            if not rect:
                continue
            bbox = ap.get("/BBox", [0, 0, rect[2]-rect[0], rect[3]-rect[1]])
            resources = page["/Resources"]
            if "/XObject" not in resources:
                resources[NameObject("/XObject")] = DictionaryObject()
            name = NameObject("/BTS" + uuid.uuid4().hex)
            resources["/XObject"][name] = writer._add_object(ap)
            sx = (float(rect[2])-float(rect[0])) / (float(bbox[2])-float(bbox[0]) or 1)
            sy = (float(rect[3])-float(rect[1])) / (float(bbox[3])-float(bbox[1]) or 1)
            content = DecodedStreamObject()
            existing_stream = page.get_contents()
            existing = existing_stream.get_data() if existing_stream is not None else b""
            content.set_data(existing + f"\nq {sx} 0 0 {sy} {float(rect[0])-float(bbox[0])*sx} {float(rect[1])-float(bbox[1])*sy} cm {name} Do Q\n".encode())
            page[NameObject("/Contents")] = writer._add_object(content)
        page.pop("/Annots", None)
    writer._root_object.pop("/AcroForm", None)
    page = writer.pages[1]
    packet = io.BytesIO(); overlay = canvas.Canvas(packet, pagesize=(float(page.mediabox.width),float(page.mediabox.height)))
    overlay.drawImage(ImageReader(io.BytesIO(assets["stamp"])), 24, 283, width=115, height=48, preserveAspectRatio=True, mask="auto")
    overlay.drawImage(ImageReader(io.BytesIO(assets["signature"])), 137, 290, width=65, height=38, preserveAspectRatio=True, mask="auto")
    overlay.save(); packet.seek(0); page.merge_page(PdfReader(packet).pages[0])
    output=io.BytesIO(); writer.write(output)
    roles = [("employer", 38), ("apprentice", 205)] + ([("guardian", 373)] if needs_guardian(values) else [])
    return output.getvalue(), [{"role": role, "type": "signature", "page": 2, "x": x, "y": 671,
                                "width": 150, "height": 45} for role,x in roles]


def generate_documents(store, record_id, values, settings, assets, existing=None):
    kind_errors = {kind: document_errors(kind, values, settings) for kind in ("formation", "mobilite")
                   if kind != "mobilite" or settings.get("teaching_mode") == "presentiel"}
    errors = list(dict.fromkeys(item for items in kind_errors.values() for item in items))
    if errors:
        raise WorkspaceError("À compléter pour générer les conventions : " + "; ".join(errors))
    package_id = uuid.uuid4().hex
    directory = store.path(package_id); directory.mkdir(mode=0o700)
    package = {"id": package_id, "created_at": now(), "fingerprint": fingerprint(values, settings),
               "values": dict(values), "settings": dict(settings), "documents": {}, "signature": {}, "opco": {}}
    cerfa, fields = prepare_cerfa(values, assets)
    (directory / "cerfa.pdf").write_bytes(cerfa)
    package["documents"]["cerfa"] = {"pdf": f"{package_id}/cerfa.pdf", "fields": fields, "sha256": hashlib.sha256(cerfa).hexdigest()}
    for kind in kind_errors:
        docx_path = directory / f"{kind}.docx"
        docx_path.write_bytes(fill_convention(kind, values, settings, assets))
        pdf_path = convert_pdf(docx_path)
        fields = extract_signature_fields(pdf_path)
        expected = {"employer"} if kind == "formation" else {"employer", "apprentice"} | ({"guardian"} if needs_guardian(values) else set())
        if {field["role"] for field in fields} != expected:
            raise WorkspaceError("Les zones de signature de la convention sont incomplètes.")
        package["documents"][kind] = {"pdf": f"{package_id}/{kind}.pdf", "docx": f"{package_id}/{kind}.docx",
            "fields": fields, "sha256": hashlib.sha256(pdf_path.read_bytes()).hexdigest()}
    return package
