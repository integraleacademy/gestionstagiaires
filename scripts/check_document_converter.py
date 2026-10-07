#!/usr/bin/env python3
"""Exercise the installed converter with synthetic files, without Flask or keys.

Run during the container build so a missing Office filter, image codec or Linux
worker restriction fails the build instead of failing a trainee's upload.
"""

from io import BytesIO
from pathlib import Path
import sys
import tempfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from PIL import Image
from pypdf import PdfReader

import document_conversion as conversion


def require(condition, message):
    if not condition:
        raise RuntimeError(message)


def read_pdf(data, filename, pages):
    require(filename.endswith(".pdf") and data.startswith(b"%PDF-"), "Conversion did not produce a PDF")
    reader = PdfReader(BytesIO(data), strict=True)
    require(len(reader.pages) == pages, f"Expected {pages} PDF pages, got {len(reader.pages)}")
    return reader


def check_office():
    txt, name = conversion.convert_upload("Synthetic text document\nLisibilité vérifiée.\n".encode(), "example.txt")
    reader = read_pdf(txt, name, 1)
    text = reader.pages[0].extract_text() or ""
    require("Synthetic text document" in text and "Lisibilité" in text, "Text import lost content or UTF-8 accents")

    source = Document()
    source.add_paragraph("SYNTHETIC FIRST PAGE")
    source.add_paragraph("Document d’essai, sans données personnelles.")
    source.add_page_break()
    source.add_paragraph("SYNTHETIC SECOND PAGE")
    incoming = BytesIO()
    source.save(incoming)
    output, name = conversion.convert_upload(incoming.getvalue(), "two-pages.docx")
    reader = read_pdf(output, name, 2)
    require("SYNTHETIC FIRST PAGE" in (reader.pages[0].extract_text() or ""), "DOCX first page missing")
    require("SYNTHETIC SECOND PAGE" in (reader.pages[1].extract_text() or ""), "DOCX second page missing")
    unchanged, _ = conversion.convert_upload(output, "existing.pdf")
    require(unchanged == output, "Existing PDF was modified")
    print("OK: real TXT and two-page DOCX conversions; existing PDF preserved", flush=True)


def check_images():
    with Image.new("RGB", (64, 96), "white") as sample:
        for fmt, suffix in (("JPEG", "jpg"), ("PNG", "png")):
            encoded = BytesIO()
            sample.save(encoded, format=fmt)
            original = encoded.getvalue()
            output, name = conversion.convert_upload(original, "synthetic." + suffix)
            read_pdf(output, name, 1)
            portrait, photo_name = conversion.convert_upload(original, "portrait." + suffix, photo=True)
            require(portrait == original and photo_name.endswith("." + suffix), fmt + " identity image was modified")

        # Exercise the real HEVC encoder and decoder supplied by pillow-heif.
        # A registry-only check would miss a wheel without a working codec.
        heic = BytesIO()
        sample.save(heic, format="HEIF")
        output, name = conversion.convert_upload(heic.getvalue(), "synthetic.heic")
        read_pdf(output, name, 1)
        portrait, photo_name = conversion.convert_upload(heic.getvalue(), "portrait.heic", photo=True)
        with Image.open(BytesIO(portrait)) as decoded:
            decoded.load()
            require(decoded.format == "JPEG" and decoded.size == sample.size and photo_name.endswith(".jpg"), "HEIC portrait did not remain an image")

        second = Image.new("RGB", sample.size, "black")
        try:
            multipage = BytesIO()
            sample.save(multipage, format="TIFF", save_all=True, append_images=[second])
            output, name = conversion.convert_upload(multipage.getvalue(), "two-pages.tiff")
            read_pdf(output, name, 2)
        finally:
            second.close()
    print("OK: JPEG, PNG, HEIC and two-page TIFF; portraits remain images", flush=True)


def check_linux_worker():
    if not sys.platform.startswith("linux"):
        print("SKIP: Linux socket isolation probe (not a Linux host)", flush=True)
        return
    # No connections are attempted. Creating an Internet socket must itself be
    # denied, through the same executable, environment and worker as Office.
    probe = """
import errno
import socket
for family in (socket.AF_INET, socket.AF_INET6):
    try:
        candidate = socket.socket(family, socket.SOCK_STREAM)
    except OSError as error:
        if error.errno != errno.EPERM:
            raise
    else:
        candidate.close()
        raise RuntimeError('Internet sockets are not blocked')
"""
    with tempfile.TemporaryDirectory(prefix="converter-worker-check-") as directory:
        conversion._run_office([sys.executable, "-c", probe], Path(directory))
    print("OK: Linux worker blocks IPv4 and IPv6 sockets", flush=True)


def main():
    check_linux_worker()
    check_images()
    check_office()
    print("Document converter checks passed.", flush=True)


if __name__ == "__main__":
    main()
